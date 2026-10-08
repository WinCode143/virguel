"""Área de Contabilidad: egresos, comprobantes, reintegros, presupuestos y costos fijos.

Acceso: gerencia y el rol Contabilidad.
"""
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django import forms
from django.contrib import messages
from django.db.models import Count, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from core.notificaciones import notificar
from core.roles import CONTABILIDAD, GERENCIA, requiere_rol
from operaciones.models import GastoOrden

from .models import CategoriaEgreso, CostoFijo, Egreso, Presupuesto
from .proyeccion import inicio_mes, proyectar, sumar_meses

ROLES = (GERENCIA, CONTABILIDAD)


class EgresoForm(forms.ModelForm):
    class Meta:
        model = Egreso
        fields = ["fecha", "categoria", "monto", "descripcion", "proveedor", "numero_comprobante", "comprobante"]
        widgets = {"fecha": forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d")}


def _mes(request):
    try:
        return inicio_mes(date.fromisoformat(request.GET.get("mes", "") + "-01"))
    except ValueError:
        return inicio_mes(timezone.localdate())


def presupuesto_vs_real(mes):
    """Por categoría: presupuesto, gastado y semáforo (verde ≤ presupuesto; amarillo ≤ +tolerancia; rojo >)."""
    fin = sumar_meses(mes, 1) - timedelta(days=1)
    real = dict(Egreso.objects.filter(fecha__range=(mes, fin)).values_list("categoria").annotate(t=Sum("monto")))
    pres = {p.categoria_id: p.monto for p in Presupuesto.objects.filter(mes=mes)}
    filas = []
    for c in CategoriaEgreso.objects.all():
        r, p = real.get(c.id, Decimal("0")), pres.get(c.id)
        if p is None and not r:
            continue
        if p is None:
            estado, uso = "info", None
        else:
            uso = float(r / p * 100) if p else None
            estado = "ok" if r <= p else "aviso" if r <= p * (1 + Decimal(c.tolerancia_presupuesto) / 100) else "critico"
        filas.append({"c": c, "presupuesto": p, "real": r, "uso": uso, "estado": estado,
                      "diferencia": (p - r) if p is not None else None})
    return filas


@requiere_rol(*ROLES)
def panel(request):
    mes = _mes(request)
    fin = sumar_meses(mes, 1) - timedelta(days=1)
    filas = presupuesto_vs_real(mes)
    total_pres = sum((f["presupuesto"] or 0) for f in filas)
    total_real = sum(f["real"] for f in filas)
    from tablero.views import grafico
    con_pres = [f for f in filas if f["presupuesto"]]
    g = grafico("bar", [f["c"].nombre for f in con_pres], [
        {"nombre": "Presupuesto", "datos": [f["presupuesto"] for f in con_pres], "serie": 1},
        {"nombre": "Gastado", "datos": [f["real"] for f in con_pres], "colores": [f"--{f['estado']}" for f in con_pres]},
    ], formato="pesos", horizontal=True)
    return render(request, "finanzas/panel.html", {
        "mes": mes, "anterior": sumar_meses(mes, -1), "siguiente": sumar_meses(mes, 1), "filas": filas, "g": g,
        "total_pres": total_pres, "total_real": total_real,
        "gastos_pendientes": GastoOrden.objects.filter(estado="pendiente").count(),
        "sin_comprobante": Egreso.objects.filter(fecha__range=(mes, fin), automatico=False, comprobante="").count(),
        "proy": proyectar(3)})


@requiere_rol(*ROLES)
def egresos(request):
    mes = _mes(request)
    fin = sumar_meses(mes, 1) - timedelta(days=1)
    qs = Egreso.objects.filter(fecha__range=(mes, fin)).select_related("categoria", "cargado_por")
    cat, origen, q = request.GET.get("categoria", ""), request.GET.get("origen", ""), request.GET.get("q", "").strip()
    if cat:
        qs = qs.filter(categoria_id=cat)
    if origen in ("manual", "automatico"):
        qs = qs.filter(automatico=origen == "automatico")
    if q:
        qs = qs.filter(descripcion__icontains=q) | qs.filter(proveedor__icontains=q)
    if request.GET.get("formato") == "excel":
        return _excel(qs, mes)
    return render(request, "finanzas/egresos.html", {
        "egresos": qs.order_by("-fecha", "-id")[:500], "total": qs.aggregate(t=Sum("monto"))["t"] or 0,
        "mes": mes, "anterior": sumar_meses(mes, -1), "siguiente": sumar_meses(mes, 1),
        "categorias": CategoriaEgreso.objects.all(), "cat": cat, "origen": origen, "q": q})


@requiere_rol(*ROLES)
def egreso_form(request, pk=None):
    e = get_object_or_404(Egreso, pk=pk, automatico=False) if pk else None
    if request.method == "POST":
        if e and request.POST.get("accion") == "eliminar":
            e.delete()
            messages.success(request, "Egreso eliminado.")
            return redirect("finanzas:egresos")
        f = EgresoForm(request.POST, request.FILES, instance=e)
        if f.is_valid():
            nuevo = f.save(commit=False)
            if not nuevo.cargado_por_id:
                nuevo.cargado_por = request.user
            nuevo.save()
            messages.success(request, "Egreso guardado.")
            return redirect(f"{request.path if 'otro' in request.POST else '/finanzas/egresos/'}")
    else:
        f = EgresoForm(instance=e)
    return render(request, "finanzas/egreso_form.html", {"form": f, "e": e})


@requiere_rol(*ROLES)
def gastos_campo(request):
    """Revisión de gastos que cargan los técnicos al cerrar llamadas (reintegros)."""
    if request.method == "POST":
        g = get_object_or_404(GastoOrden, pk=request.POST.get("gasto"))
        g.estado = "aprobado" if request.POST.get("accion") == "aprobar" else "rechazado"
        g.save()  # el egreso se actualiza solo (rechazado = se quita)
        notificar(g.tecnico, f"Gasto {'aprobado' if g.estado == 'aprobado' else 'rechazado'}: {g.descripcion}",
                  f"${g.monto:,.0f}" + (f" · {request.POST.get('motivo')}" if request.POST.get("motivo") else ""), "/app/")
        messages.success(request, f"Gasto {g.get_estado_display().lower()}.")
        return redirect("finanzas:gastos_campo")
    estado = request.GET.get("estado", "pendiente")
    qs = GastoOrden.objects.select_related("tecnico", "orden", "orden__tipo").order_by("-fecha")
    if estado:
        qs = qs.filter(estado=estado)
    por_tecnico = (GastoOrden.objects.filter(estado="aprobado", fecha__gte=inicio_mes(timezone.localdate()))
                   .values("tecnico__apellido", "tecnico__nombre").annotate(t=Sum("monto"), n=Count("id")).order_by("-t"))
    return render(request, "finanzas/gastos_campo.html", {"gastos": qs[:200], "estado": estado,
                                                           "por_tecnico": por_tecnico})


def _dec(v):
    v = (v or "").strip().replace("$", "").replace(" ", "")
    if not v:
        return None
    if "," in v:
        v = v.replace(".", "").replace(",", ".")
    try:
        return Decimal(v)
    except InvalidOperation:
        return False


@requiere_rol(*ROLES)
def presupuestos(request):
    mes = _mes(request)
    meses = [sumar_meses(mes, i) for i in range(3)]
    cats = list(CategoriaEgreso.objects.all())
    if request.method == "POST":
        if request.POST.get("accion") == "copiar":
            anterior = sumar_meses(mes, -1)
            n = 0
            for p in Presupuesto.objects.filter(mes=anterior):
                _, creado = Presupuesto.objects.get_or_create(categoria=p.categoria, mes=mes, defaults={"monto": p.monto})
                n += creado
            messages.success(request, f"Copiados {n} presupuestos de {anterior:%m/%Y}.")
        else:
            errores = 0
            for c in cats:
                for m in meses:
                    v = _dec(request.POST.get(f"p_{c.id}_{m:%Y%m}"))
                    if v is False:
                        errores += 1
                    elif v is None:
                        Presupuesto.objects.filter(categoria=c, mes=m).delete()
                    else:
                        Presupuesto.objects.update_or_create(categoria=c, mes=m, defaults={"monto": v})
            messages.error(request, f"{errores} valor(es) no son números.") if errores else messages.success(
                request, "Presupuestos guardados.")
        return redirect(f"{request.path}?mes={mes:%Y-%m}")
    actuales = {(p.categoria_id, p.mes): p.monto for p in Presupuesto.objects.filter(mes__in=meses)}
    filas = [{"c": c, "valores": [(m, actuales.get((c.id, m))) for m in meses]} for c in cats]
    totales = [sum((actuales.get((c.id, m)) or 0) for c in cats) for m in meses]
    return render(request, "finanzas/presupuestos.html", {"filas": filas, "meses": meses, "mes": mes, "totales": totales,
                                                           "anterior": sumar_meses(mes, -1)})


@requiere_rol(*ROLES)
def configuracion(request):
    if request.method == "POST":
        accion = request.POST.get("accion")
        if accion == "categoria":
            from django.utils.text import slugify
            nombre = request.POST.get("nombre", "").strip()
            if nombre:
                CategoriaEgreso.objects.get_or_create(nombre=nombre, defaults={"codigo": slugify(nombre)[:30]})
                messages.success(request, f"Categoría «{nombre}» creada.")
        elif accion == "tolerancia":
            c = get_object_or_404(CategoriaEgreso, pk=request.POST.get("categoria"))
            v = request.POST.get("tolerancia", "")
            if v.isdigit():
                c.tolerancia_presupuesto = int(v)
                c.save(update_fields=["tolerancia_presupuesto"])
        elif accion == "costo":
            monto = _dec(request.POST.get("monto"))
            cat = CategoriaEgreso.objects.filter(pk=request.POST.get("categoria")).first()
            if cat and monto:
                CostoFijo.objects.create(categoria=cat, descripcion=request.POST.get("descripcion", "")[:150], monto_mensual=monto)
                messages.success(request, "Costo fijo agregado.")
        elif accion in ("costo_editar", "costo_baja"):
            cf = get_object_or_404(CostoFijo, pk=request.POST.get("costo"))
            if accion == "costo_baja":
                cf.activo = not cf.activo
            else:
                monto = _dec(request.POST.get("monto"))
                if monto:
                    cf.monto_mensual = monto
            cf.save()
            messages.success(request, "Costo fijo actualizado.")
        return redirect("finanzas:configuracion")
    return render(request, "finanzas/configuracion.html", {
        "categorias": CategoriaEgreso.objects.annotate(n=Count("egresos")),
        "costos": CostoFijo.objects.select_related("categoria").order_by("-activo", "descripcion")})


def _excel(qs, mes):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = f"Egresos {mes:%m-%Y}"
    cab = ["Fecha", "Categoría", "Descripción", "Proveedor", "N° comprobante", "Monto", "Origen", "Cargado por",
           "Tiene comprobante"]
    ws.append(cab)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2A78D6")
    for e in qs.order_by("fecha", "id"):
        ws.append([e.fecha, e.categoria.nombre, e.descripcion, e.proveedor, e.numero_comprobante, float(e.monto),
                   "Automático" if e.automatico else "Manual",
                   (e.cargado_por.get_full_name() or e.cargado_por.username) if e.cargado_por else "",
                   "Sí" if e.comprobante else "No"])
    ws.append([])
    ws.append(["", "", "", "", "TOTAL", f"=SUM(F2:F{ws.max_row - 1})"])
    for col, ancho in zip("ABCDEFGHI", (12, 22, 50, 22, 16, 14, 12, 18, 10)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in fila:
            c.number_format = '"$"#,##0.00'
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="egresos_{mes:%Y_%m}.xlsx"'
    return resp
