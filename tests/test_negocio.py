"""Tests de las reglas de negocio principales.

Ejecutar:  python manage.py test tests
"""
import io
from datetime import timedelta
from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.utils import timezone

from capacitacion.evaluacion import Diagnostico, evaluar_tecnicos
from capacitacion.models import Capacitacion, Curso, Participacion
from core.models import Alerta, Cliente, Parametros, Persona, Zona
from core.services import SincronizadorAlertas
from finanzas.models import Egreso
from flota.models import ServiceRealizado, TipoService, Vehiculo, proximos_services
from incidentes.models import Siniestro
from inventario.models import LoteIngreso, Material, Salida
from inventario.services import StockInsuficiente, lotes_envejecidos, registrar_salida, stock_diario
from operaciones.models import Jornada, OrdenTrabajo, TipoTarea
from supervision.evaluacion import evaluar_supervisores
from supervision.models import AccionCorrectiva, EncuestaSupervisor, InformeControl
from tablero.planificacion import capacidad, probabilidad_decodificador

HOY = timezone.localdate()


# En los tests, los usuarios nuevos ya aceptaron el aviso de privacidad (salvo que el test lo cambie).
def _aceptar_privacidad(sender, instance, created, **kwargs):
    if created:
        from core.models import CuentaUsuario
        CuentaUsuario.objects.get_or_create(usuario=instance, defaults={"acepto_privacidad": timezone.now()})


from django.db.models.signals import post_save  # noqa: E402
post_save.connect(_aceptar_privacidad, sender=User, dispatch_uid="tests_privacidad")


def persona(legajo, rol="tecnico", **kw):
    return Persona.objects.create(legajo=legajo, nombre=legajo, apellido=legajo, rol=rol,
                                  fecha_ingreso=kw.pop("fecha_ingreso", HOY - timedelta(days=400)), **kw)


class StockTests(TestCase):
    def setUp(self):
        self.m = Material.objects.create(codigo="X", nombre="Cable", costo_unitario=Decimal("10"))

    def test_salida_fifo_consume_primero_lo_mas_viejo(self):
        viejo = LoteIngreso.objects.create(material=self.m, fecha=HOY - timedelta(days=30), cantidad=10,
                                           costo_unitario=Decimal("5"))
        nuevo = LoteIngreso.objects.create(material=self.m, fecha=HOY, cantidad=10, costo_unitario=Decimal("8"))
        s = registrar_salida(Salida(material=self.m, cantidad=Decimal("14")))
        viejo.refresh_from_db(); nuevo.refresh_from_db()
        self.assertEqual(viejo.cantidad_disponible, 0)
        self.assertEqual(nuevo.cantidad_disponible, 6)
        self.assertEqual(s.costo_total, Decimal("10") * 5 + Decimal("4") * 8)

    def test_salida_sin_stock_falla_salvo_que_se_permita(self):
        LoteIngreso.objects.create(material=self.m, cantidad=2)
        with self.assertRaises(StockInsuficiente):
            registrar_salida(Salida(material=self.m, cantidad=5))
        s = registrar_salida(Salida(material=self.m, cantidad=5), permitir_negativo=True)
        self.assertEqual(s.costo_total, Decimal("50"))

    def test_antiguedad_marca_lotes_de_mas_de_60_dias(self):
        LoteIngreso.objects.create(material=self.m, fecha=HOY - timedelta(days=61), cantidad=3)
        LoteIngreso.objects.create(material=self.m, fecha=HOY - timedelta(days=50), cantidad=3)
        LoteIngreso.objects.create(material=self.m, fecha=HOY - timedelta(days=5), cantidad=3)
        estados = [f["estado"] for f in lotes_envejecidos(HOY)]
        self.assertEqual(estados, ["critico", "aviso", "ok"])

    def test_stock_diario_cruza_con_tecnicos_en_calle(self):
        LoteIngreso.objects.create(material=self.m, fecha=HOY - timedelta(days=40), cantidad=100)
        tecs = [persona(f"T{i}") for i in range(4)]
        for d in range(1, 11):
            for t in tecs:
                Jornada.objects.create(fecha=HOY - timedelta(days=d), tecnico=t)
            registrar_salida(Salida(material=self.m, cantidad=Decimal("4"), fecha=HOY - timedelta(days=d)))
        for t in tecs:
            Jornada.objects.create(fecha=HOY, tecnico=t)
        fila = stock_diario(HOY)["filas"][0]
        self.assertEqual(stock_diario(HOY)["tecnicos_en_calle"], 4)
        self.assertEqual(fila["consumo_tecnico_dia"], Decimal("1"))       # 40 unidades / 40 jornadas
        self.assertEqual(fila["stock"], Decimal("60"))
        self.assertAlmostEqual(float(fila["dias_cobertura"]), 15.0)    # 60 / (1 x 4 técnicos)


class PlanificacionTests(TestCase):
    def test_capacidad_sin_hectareas_es_tecnicos_por_clientes(self):
        cap = capacidad(10, prob_min=0.5, prob_max=0.6)
        self.assertEqual((cap.clientes_min, cap.clientes_max), (60, 60))
        self.assertEqual((cap.deco_min, cap.deco_max), (30, 36))

    def test_capacidad_hectareas_y_clientes(self):
        par = Parametros.actual()  # 5-6 ha por cuadrilla, 2 técnicos por cuadrilla, 6 clientes/técnico
        par.usar_hectareas = True
        par.save()
        cap = capacidad(10, densidad=1.0, prob_min=0.5, prob_max=0.6)
        self.assertEqual(cap.cuadrillas, 5)
        self.assertEqual((cap.ha_min, cap.ha_max), (25, 30))
        self.assertEqual((cap.clientes_min, cap.clientes_max), (25, 30))
        self.assertEqual((cap.deco_min, cap.deco_max), (12.5, 18))

    def test_capacidad_limitada_por_tope_de_visitas(self):
        par = Parametros.actual()
        par.usar_hectareas = True
        par.save()
        cap = capacidad(10, densidad=5.0)
        self.assertEqual(cap.clientes_max, 60)  # 10 técnicos x 6 visitas

    def test_probabilidad_sin_datos_usa_supuesto(self):
        est = probabilidad_decodificador(HOY)
        self.assertAlmostEqual(est.media, 0.55)

    def test_probabilidad_se_corrige_con_datos_reales(self):
        tipo = TipoTarea.objects.create(codigo="I", nombre="Inst", puede_requerir_decodificador=True)
        for i in range(400):
            OrdenTrabajo.objects.create(numero=str(i), tipo=tipo, estado="completada", fecha_ejecucion=HOY,
                                        decodificador_solicitado=i < 300)  # 75% real
        est = probabilidad_decodificador(HOY)
        self.assertGreater(est.media, 0.72)
        self.assertLess(est.inf90, est.media)
        self.assertGreater(est.sup90, est.media)


class EvaluacionTecnicosTests(TestCase):
    """Escenario: 6 técnicos normales, uno lento pero prolijo y uno lento y problemático."""

    def setUp(self):
        self.tipo = TipoTarea.objects.create(codigo="R", nombre="Rep", minutos_estandar=60)
        self.sup = persona("S1", rol="supervisor")
        self.normales = [persona(f"N{i}", supervisor=self.sup) for i in range(6)]
        self.lento = persona("LENTO", supervisor=self.sup)
        self.malo = persona("MALO", supervisor=self.sup)
        n = 0
        for d in range(1, 41):
            f = HOY - timedelta(days=d)
            for t in self.normales + [self.lento, self.malo]:
                Jornada.objects.create(fecha=f, tecnico=t)
                cant = 3 if t in (self.lento, self.malo) else 5
                for k in range(cant):
                    n += 1
                    o = OrdenTrabajo.objects.create(numero=str(n), tipo=self.tipo, tecnico=t, estado="completada",
                                                    fecha_programada=f, fecha_ejecucion=f, minutos_reales=60)
                    if t is self.malo and k == 0:  # un retrabajo por día: 1 de cada 3 trabajos
                        n += 1
                        OrdenTrabajo.objects.create(numero=str(n), tipo=self.tipo, es_retrabajo=True,
                                                    orden_original=o, fecha_programada=f, estado="completada")
        for k in range(4):
            AccionCorrectiva.objects.create(tecnico=self.malo, tipo="apercibimiento", fecha=HOY - timedelta(days=k * 7))
        Siniestro.objects.create(numero="S1", tipo="cano_pinchado", gravedad="grave", tecnico=self.malo,
                                 descripcion="x", fecha=HOY - timedelta(days=10))

    def diag(self):
        return {e.tecnico.legajo: e for e in evaluar_tecnicos(HOY, dias=40)}

    def test_lento_pero_prolijo_necesita_capacitacion(self):
        self.assertEqual(self.diag()["LENTO"].diagnostico, Diagnostico.CAPACITACION)

    def test_lento_y_problematico_es_riesgo_alto(self):
        e = self.diag()["MALO"]
        self.assertEqual(e.diagnostico, Diagnostico.RIESGO)
        self.assertGreater(e.riesgo, self.diag()["LENTO"].riesgo)

    def test_normales_adecuados(self):
        d = self.diag()
        self.assertTrue(all(d[t.legajo].diagnostico == Diagnostico.ADECUADO for t in self.normales))

    def test_capacitado_y_mejorando(self):
        # el lento recibió capacitación en producción y en los últimos días rinde como el resto
        curso = Curso.objects.create(nombre="Acompañamiento")
        cap = Capacitacion.objects.create(curso=curso, fecha=HOY - timedelta(days=15), en_produccion=True)
        Participacion.objects.create(capacitacion=cap, persona=self.lento)
        for d in range(1, 14):
            f = HOY - timedelta(days=d)
            for i in range(2):
                OrdenTrabajo.objects.create(numero=f"X{d}-{i}", tipo=self.tipo, tecnico=self.lento,
                                            estado="completada", fecha_programada=f, fecha_ejecucion=f)
        e = self.diag()["LENTO"]
        self.assertGreater(e.tendencia, 10)
        self.assertIn(e.diagnostico, (Diagnostico.MEJORANDO, Diagnostico.ADECUADO))

    def test_evaluar_una_sola_persona_usa_mediana_de_todo_el_plantel(self):
        sola = evaluar_tecnicos(HOY, dias=40, tecnicos=[self.malo])[0]
        self.assertAlmostEqual(sola.indice_productividad, self.diag()["MALO"].indice_productividad)
        self.assertLess(sola.indice_productividad, 0.7)
        self.assertEqual(sola.diagnostico, Diagnostico.RIESGO)

    def test_ingresante_en_curva_de_aprendizaje(self):
        self.lento.fecha_ingreso = HOY - timedelta(days=30)
        self.lento.save()
        self.assertEqual(self.diag()["LENTO"].diagnostico, Diagnostico.APRENDIZAJE)


class SupervisoresTests(TestCase):
    def test_encuesta_y_desvios_afectan_puntaje(self):
        bueno, malo = persona("S1", rol="supervisor"), persona("S2", rol="supervisor")
        tb, tm = persona("T1", supervisor=bueno), persona("T2", supervisor=malo)
        for d in range(1, 11):
            f = HOY - timedelta(days=d)
            for t, s, nota in ((tb, bueno, 5), (tm, malo, 2)):
                Jornada.objects.create(fecha=f, tecnico=t)
                EncuestaSupervisor.objects.create(fecha=f, tecnico=t, supervisor=s, respondida=timezone.now(),
                                                  trato=nota, claridad=nota, apoyo=nota, presencia=nota)
            inf = InformeControl.objects.create(supervisor=bueno, tecnico=tb, fecha=f, puntaje=2,
                                                desvio_detectado=True, descripcion="x" * 100)
            AccionCorrectiva.objects.create(informe=inf, tecnico=tb, tipo="recapacitacion", fecha=f + timedelta(days=1))
            InformeControl.objects.create(supervisor=malo, tecnico=tm, fecha=f, puntaje=2, desvio_detectado=True)
        evs = {e.supervisor.legajo: e for e in evaluar_supervisores(HOY, 30)}
        self.assertEqual(evs["S1"].nota_encuesta, 5)
        self.assertEqual(evs["S1"].desvios_con_accion, 10)
        self.assertEqual(evs["S2"].desvios_con_accion, 0)
        self.assertGreater(evs["S1"].score_total, evs["S2"].score_total + 30)


class AutomatismosTests(TestCase):
    def test_egresos_automaticos(self):
        m = Material.objects.create(codigo="X", nombre="Cable", costo_unitario=Decimal("10"))
        lote = LoteIngreso.objects.create(material=m, cantidad=5)
        self.assertEqual(Egreso.objects.get(origen=f"lote:{lote.id}").monto, Decimal("50"))
        s = Siniestro.objects.create(numero="S1", tipo="otro", descripcion="x")
        self.assertFalse(Egreso.objects.filter(origen=f"siniestro:{s.id}").exists())  # sin costo real aún
        s.costo_real = Decimal("1000"); s.save()
        self.assertEqual(Egreso.objects.get(origen=f"siniestro:{s.id}").monto, Decimal("1000"))
        lote.delete()
        self.assertFalse(Egreso.objects.filter(origen=f"lote:{lote.id}").exists())

    def test_alertas_se_abren_y_cierran_solas(self):
        with SincronizadorAlertas("stock") as s:
            s.alerta("a", "A"); s.alerta("b", "B")
        with SincronizadorAlertas("stock") as s:
            s.alerta("a", "A actualizada")
        abiertas = Alerta.objects.filter(resuelta=False)
        self.assertEqual([a.titulo for a in abiertas], ["A actualizada"])

    def test_service_vencido_por_km(self):
        v = Vehiculo.objects.create(patente="AA1", marca="X", modelo="Y", anio=2020, km_actual=21000)
        ts = TipoService.objects.create(nombre="Aceite", cada_km=10000, cada_dias=365)
        ServiceRealizado.objects.create(vehiculo=v, tipo=ts, fecha=HOY - timedelta(days=30), km=10000)
        p = proximos_services(v, HOY)[0]
        self.assertEqual(p["estado"], "critico")
        self.assertEqual(p["km_restantes"], -1000)


class AccesoTests(TestCase):
    def setUp(self):
        self.ger = User.objects.create_user("ger", password="x")
        self.ger.groups.add(Group.objects.create(name="Gerencia"))
        self.sup_p = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="x"))
        self.tec_p = persona("T1", supervisor=self.sup_p, usuario=User.objects.create_user("tec", password="x"))

    def test_cada_rol_llega_a_su_version(self):
        for user, destino in (("ger", "/tablero/"), ("sup", "/tablero/"), ("tec", "/app/")):
            self.client.login(username=user, password="x")
            self.assertRedirects(self.client.get("/"), destino, fetch_redirect_response=False)
        self.client.login(username="sup", password="x")
        r = self.client.get("/", HTTP_USER_AGENT="Mozilla/5.0 (Linux; Android 14) Mobile")
        self.assertRedirects(r, "/app/", fetch_redirect_response=False)

    def test_permisos(self):
        self.client.login(username="tec", password="x")
        self.assertEqual(self.client.get("/tablero/").status_code, 403)
        self.client.login(username="sup", password="x")
        self.assertEqual(self.client.get("/tablero/finanzas/").status_code, 403)
        self.assertEqual(self.client.get("/tablero/supervisores/").status_code, 403)
        self.assertEqual(self.client.get("/tablero/tecnicos/").status_code, 200)
        self.assertEqual(self.client.get("/api/powerbi/ordenes.csv").status_code, 403)

    def test_encuesta_por_link_sin_login(self):
        e = EncuestaSupervisor.objects.create(fecha=HOY, tecnico=self.tec_p, supervisor=self.sup_p)
        r = self.client.post(f"/encuesta/{e.token}/", {"trato": 4, "claridad": 5, "apoyo": 3, "presencia": 4})
        self.assertEqual(r.status_code, 200)
        e.refresh_from_db()
        self.assertIsNotNone(e.respondida)
        self.assertEqual(e.promedio, 4)

    def test_tecnico_cierra_orden_y_descuenta_de_su_stock(self):
        from inventario.stock_tecnico import entregar, saldo
        tipo = TipoTarea.objects.create(codigo="I", nombre="Inst", puede_requerir_decodificador=True)
        m = Material.objects.create(codigo="C", nombre="Cable", costo_unitario=Decimal("2"))
        LoteIngreso.objects.create(material=m, cantidad=100)
        entregar(self.tec_p, m, 50)  # el depósito le entrega 50: salen del depósito
        self.assertEqual((m.stock_actual, saldo(self.tec_p, m)), (50, 50))
        ot = OrdenTrabajo.objects.create(numero="1", tipo=tipo, tecnico=self.tec_p, estado="asignada")
        self.client.login(username="tec", password="x")
        self.assertEqual(self.client.get(f"/app/orden/{ot.id}/").status_code, 200)
        r = self.client.post(f"/app/orden/{ot.id}/", {"resultado": "completada", "minutos_reales": 50,
                                                      "decodificador_solicitado": "on", "decodificadores_instalados": 1,
                                                      "material_1": m.id, "cantidad_1": "30",
                                                      "series_instaladas": "ABC123", "conforme_nombre": "Ana"})
        self.assertEqual(r.status_code, 302)
        ot.refresh_from_db()
        self.assertEqual((ot.estado, ot.series_instaladas, ot.conforme_nombre), ("completada", "ABC123", "Ana"))
        self.assertTrue(ot.decodificador_solicitado)
        self.assertEqual(saldo(self.tec_p, m), 20)      # se descontó de SU stock
        self.assertEqual(m.stock_actual, 50)            # el depósito no cambia al cerrar
        self.assertEqual(Salida.objects.get(orden=ot).costo_total, Decimal("60"))

    def test_orden_no_resuelta_exige_motivo(self):
        tipo = TipoTarea.objects.create(codigo="R", nombre="Rep")
        ot = OrdenTrabajo.objects.create(numero="2", tipo=tipo, tecnico=self.tec_p, estado="asignada")
        self.client.login(username="tec", password="x")
        r = self.client.post(f"/app/orden/{ot.id}/", {"resultado": "fallida"})
        self.assertEqual(r.status_code, 200)
        self.client.post(f"/app/orden/{ot.id}/", {"resultado": "fallida", "motivo_no_resuelto": "cliente_ausente"})
        ot.refresh_from_db()
        self.assertEqual((ot.estado, ot.motivo_no_resuelto), ("fallida", "cliente_ausente"))

    def test_empezar_trabajo_mide_el_tiempo(self):
        tipo = TipoTarea.objects.create(codigo="R", nombre="Rep")
        ot = OrdenTrabajo.objects.create(numero="3", tipo=tipo, tecnico=self.tec_p, estado="asignada")
        self.client.login(username="tec", password="x")
        hace = timezone.now() - timedelta(minutes=45)
        self.client.post(f"/app/orden/{ot.id}/", {"accion": "empezar", "_momento_cliente": hace.isoformat()})
        self.client.post(f"/app/orden/{ot.id}/", {"resultado": "completada"})
        ot.refresh_from_db()
        self.assertIn(ot.minutos_reales, (44, 45, 46))


class ImportacionTests(TestCase):
    def setUp(self):
        u = User.objects.create_user("ger", password="x")
        u.groups.add(Group.objects.create(name="Gerencia"))
        self.client.login(username="ger", password="x")

    def subir(self, tipo, contenido, nombre="datos.csv"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post("/tablero/importar/", {"tipo": tipo, "archivo": SimpleUploadedFile(
            nombre, contenido.encode("utf-8"), content_type="text/csv")})

    def test_importa_personas_con_usuario_y_supervisor(self):
        r = self.subir("personas", "legajo;nombre;apellido;rol;legajo_supervisor;usuario;dni\n"
                                   "S9;Ana;Sup;supervisor;;asup;111\n"
                                   "T9;Beto;Tec;tecnico;S9;btec;222\n")
        self.assertEqual(r.status_code, 302)
        t = Persona.objects.get(legajo="T9")
        self.assertEqual(t.supervisor.legajo, "S9")
        self.assertTrue(t.usuario.check_password("222"))

    def test_error_en_una_fila_no_guarda_nada(self):
        r = self.subir("materiales", "codigo,nombre,costo_unitario\nA,Cable,10\nB,Conector,abc\n")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "no es un número")
        self.assertFalse(Material.objects.exists())

    def test_importa_stock_con_formato_argentino(self):
        Material.objects.create(codigo="A", nombre="Cable")
        self.subir("stock", "codigo_material;fecha_ingreso;cantidad;costo_unitario\nA;01/08/2026;1.500;12,50\n")
        lote = LoteIngreso.objects.get()
        self.assertEqual((lote.cantidad, lote.costo_unitario), (Decimal("1500"), Decimal("12.50")))

    def test_importa_excel(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["patente", "marca", "modelo", "anio", "km_actual", "vencimiento_vtv"])
        wb.active.append(["ab 123 cd", "Fiat", "Fiorino", 2022, 45000, timezone.datetime(2026, 12, 1)])
        buf = io.BytesIO()
        wb.save(buf)
        self.client.post("/tablero/importar/", {"tipo": "vehiculos", "archivo": SimpleUploadedFile("v.xlsx", buf.getvalue())})
        v = Vehiculo.objects.get(patente="AB123CD")
        self.assertEqual((v.km_actual, v.vencimiento_vtv.month), (45000, 12))


class SinSenalTests(TestCase):
    def test_orden_enviada_tarde_conserva_fecha_de_carga(self):
        sup = persona("S1", rol="supervisor")
        tec = persona("T1", supervisor=sup, usuario=User.objects.create_user("tec", password="x"))
        tipo = TipoTarea.objects.create(codigo="R", nombre="Rep")
        ot = OrdenTrabajo.objects.create(numero="1", tipo=tipo, tecnico=tec, estado="asignada")
        self.client.login(username="tec", password="x")
        ayer = HOY - timedelta(days=1)
        self.client.post(f"/app/orden/{ot.id}/", {"resultado": "completada", "_fecha_cliente": ayer.isoformat()})
        ot.refresh_from_db()
        self.assertEqual(ot.fecha_ejecucion, ayer)
        # una fecha demasiado vieja o futura se ignora
        ot2 = OrdenTrabajo.objects.create(numero="2", tipo=tipo, tecnico=tec, estado="asignada")
        self.client.post(f"/app/orden/{ot2.id}/", {"resultado": "completada", "_fecha_cliente": "2020-01-01"})
        ot2.refresh_from_db()
        self.assertEqual(ot2.fecha_ejecucion, HOY)


class SeguridadTests(TestCase):
    def test_resolver_alerta_no_redirige_a_sitios_externos(self):
        u = User.objects.create_user("ger", password="x")
        u.groups.add(Group.objects.create(name="Gerencia"))
        self.client.login(username="ger", password="x")
        a = Alerta.objects.create(modulo="stock", clave="x", titulo="X")
        r = self.client.post(f"/tablero/alertas/{a.id}/resolver/", {"next": "https://sitio-falso.com/"})
        self.assertEqual(r["Location"], "/tablero/alertas/")
        a2 = Alerta.objects.create(modulo="stock", clave="y", titulo="Y")
        r = self.client.post(f"/tablero/alertas/{a2.id}/resolver/", {"next": "/tablero/alertas/?modulo=stock"})
        self.assertEqual(r["Location"], "/tablero/alertas/?modulo=stock")

    def test_tecnico_no_puede_cerrar_orden_ajena(self):
        a = persona("T1", usuario=User.objects.create_user("a", password="x"))
        b = persona("T2")
        tipo = TipoTarea.objects.create(codigo="R", nombre="Rep")
        ot = OrdenTrabajo.objects.create(numero="1", tipo=tipo, tecnico=b, estado="asignada")
        self.client.login(username="a", password="x")
        self.assertEqual(self.client.post(f"/app/orden/{ot.id}/", {"resultado": "completada"}).status_code, 404)
        self.assertNotEqual(a.id, b.id)


class AsignacionTests(TestCase):
    def test_reparte_por_capacidad_real_y_zona(self):
        from tablero.asignacion import proponer
        norte, sur = Zona.objects.create(nombre="N"), Zona.objects.create(nombre="S")
        rapido = persona("R", zona=norte)
        lento = persona("L", zona=norte)
        otro = persona("O", zona=sur)
        tipo = TipoTarea.objects.create(codigo="I", nombre="Inst")
        n = 0
        for d in range(1, 11):
            f = HOY - timedelta(days=d)
            for t, cant in ((rapido, 6), (lento, 2), (otro, 4)):
                Jornada.objects.create(fecha=f, tecnico=t)
                for _ in range(cant):
                    n += 1
                    OrdenTrabajo.objects.create(numero=f"h{n}", tipo=tipo, tecnico=t, estado="completada",
                                                fecha_programada=f, fecha_ejecucion=f)
        for i in range(12):  # 12 pendientes en el norte: capacidad norte = 6 + 2 = 8
            OrdenTrabajo.objects.create(numero=f"p{i}", tipo=tipo, zona=norte, fecha_programada=HOY)
        prop = proponer(HOY, Persona.objects.filter(rol="tecnico"))
        nuevas = {c.tecnico.legajo: len(c.nuevas) for c in prop["cupos"]}
        self.assertEqual(nuevas, {"R": 6, "L": 2, "O": 4})  # los 4 que sobran van al sur
        self.assertEqual(len(prop["otra_zona"]), 4)
        self.assertEqual(prop["sin_asignar"], [])


class ControlPersonalTests(TestCase):
    def setUp(self):
        from datetime import time
        Parametros.actual()
        self.sup = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="x"))
        self.t = persona("T1", supervisor=self.sup, usuario=User.objects.create_user("tec", password="x"),
                         fecha_ingreso=HOY - timedelta(days=400))
        self.t.hora_entrada, self.t.hora_salida, self.t.trabaja_sabados = time(8, 0), time(17, 0), False
        self.t.save()

    def momento(self, d, h, m=0):
        from datetime import datetime
        return timezone.make_aware(datetime(d.year, d.month, d.day, h, m))

    def lunes_pasado(self):
        return HOY - timedelta(days=HOY.weekday() + 7)

    def test_llegada_tarde_y_horas_extra(self):
        from personal.models import Asistencia
        d = self.lunes_pasado()
        a = Asistencia.objects.create(persona=self.t, fecha=d, entrada=self.momento(d, 8, 25), salida=self.momento(d, 19, 25))
        self.assertEqual(a.minutos_tarde, 25)
        self.assertEqual(a.horas_trabajadas, Decimal("11.00"))
        self.assertEqual(a.horas_extra, Decimal("2.00"))  # jornada normal de 9 h
        b = Asistencia.objects.create(persona=self.t, fecha=d + timedelta(days=1), entrada=self.momento(d, 8, 9))
        self.assertEqual(b.minutos_tarde, 0)  # dentro de la tolerancia de 10 min

    def test_resumen_distingue_justificadas_injustificadas_y_vacaciones(self):
        from personal.indicadores import resumen
        from personal.models import Asistencia, Feriado, Novedad
        lunes = self.lunes_pasado()
        # semana lun–vie: lunes trabajó, martes enfermo, miércoles feriado, jueves vacaciones, viernes falta sin aviso
        Asistencia.objects.create(persona=self.t, fecha=lunes, entrada=self.momento(lunes, 8))
        Novedad.objects.create(persona=self.t, tipo="enfermedad", desde=lunes + timedelta(days=1),
                               hasta=lunes + timedelta(days=1), estado="aprobada")
        Feriado.objects.create(fecha=lunes + timedelta(days=2), nombre="Feriado")
        Novedad.objects.create(persona=self.t, tipo="vacaciones", desde=lunes + timedelta(days=3),
                               hasta=lunes + timedelta(days=3), estado="aprobada")
        r = resumen(lunes, lunes + timedelta(days=4), [self.t])[0]
        self.assertEqual((r.esperados, r.presentes, r.justificadas, r.injustificadas, r.no_computables),
                         (4, 1, 2, 1, 1))
        self.assertAlmostEqual(r.presentismo, 1 / 3)  # 1 presente de 3 días computables

    def test_sin_control_de_asistencia_no_hay_faltas(self):
        from personal.indicadores import resumen
        r = resumen(HOY - timedelta(days=30), HOY, [self.t])[0]
        self.assertEqual((r.esperados, r.injustificadas), (0, 0))

    def test_fichar_desde_el_celular_con_hora_del_celular(self):
        from personal.models import Asistencia
        self.client.login(username="tec", password="x")
        hace_un_rato = timezone.now() - timedelta(hours=2)
        self.client.post("/app/jornada/iniciar/", {"lat": "-34.6", "lng": "-58.4",
                                                   "_momento_cliente": hace_un_rato.isoformat()})
        a = Asistencia.objects.get(persona=self.t)
        self.assertEqual(a.entrada, hace_un_rato)
        self.assertTrue(Jornada.objects.filter(tecnico=self.t, en_calle=True).exists())
        self.client.post("/app/jornada/finalizar/", {})
        a.refresh_from_db()
        self.assertIsNotNone(a.salida)

    def test_supervisor_aprueba_solo_avisos_de_su_equipo(self):
        from personal.models import Novedad
        ajeno = persona("T2")
        n1 = Novedad.objects.create(persona=self.t, tipo="enfermedad")
        n2 = Novedad.objects.create(persona=ajeno, tipo="enfermedad")
        self.client.login(username="sup", password="x")
        self.client.post("/app/novedades/", {"novedad": n1.id, "accion": "aprobar"})
        self.assertEqual(self.client.post("/app/novedades/", {"novedad": n2.id, "accion": "aprobar"}).status_code, 404)
        n1.refresh_from_db(); n2.refresh_from_db()
        self.assertEqual((n1.estado, n2.estado), ("aprobada", "pendiente"))
        self.assertEqual(self.client.get(f"/personal/legajos/{ajeno.id}/").status_code, 404)
        self.assertEqual(self.client.get(f"/personal/legajos/{self.t.id}/").status_code, 200)

    def test_tecnico_avisa_ausencia_con_certificado(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from personal.models import Novedad
        self.client.login(username="tec", password="x")
        r = self.client.post("/app/ausencia/", {"tipo": "enfermedad", "desde": HOY.isoformat(), "hasta": HOY.isoformat(),
                                                "certificado": SimpleUploadedFile("c.jpg", b"img", content_type="image/jpeg")})
        self.assertEqual(r.status_code, 302)
        n = Novedad.objects.get(persona=self.t)
        self.assertEqual(n.estado, "pendiente")
        self.assertTrue(n.certificado.name.endswith(".jpg"))
        n.certificado.delete()
        # no puede autoasignarse una "suspensión" o "injustificada"
        r = self.client.post("/app/ausencia/", {"tipo": "injustificada", "desde": HOY.isoformat(), "hasta": HOY.isoformat()})
        self.assertEqual(r.status_code, 200)



class PartesTecnicoTests(TestCase):
    def setUp(self):
        self.sup = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="x"))
        self.t = persona("T1", supervisor=self.sup, usuario=User.objects.create_user("tec", password="x"))
        self.ger = User.objects.create_user("ger", password="x")
        self.ger.groups.add(Group.objects.create(name="Gerencia"))
        self.m = Material.objects.create(codigo="ONT", nombre="Módem", costo_unitario=Decimal("100"))
        LoteIngreso.objects.create(material=self.m, cantidad=20)

    def test_pedido_aprobado_y_entregado_pasa_al_stock_del_tecnico(self):
        from core.models import Notificacion
        from inventario.models import PedidoMaterial
        from inventario.stock_tecnico import saldo
        self.client.login(username="tec", password="x")
        self.client.post("/app/stock/pedir/", {"material_1": self.m.id, "cantidad_1": "5", "motivo": "x"})
        ped = PedidoMaterial.objects.get()
        self.assertTrue(Notificacion.objects.filter(persona=self.sup).exists())  # se avisó al supervisor
        self.client.login(username="sup", password="x")
        self.client.post("/app/pedidos/", {"pedido": ped.id, "accion": "aprobar"})
        ped.refresh_from_db()
        self.assertEqual(ped.estado, "aprobado")
        self.client.login(username="ger", password="x")
        item = ped.items.get()
        self.client.post("/tablero/pedidos/", {"pedido": ped.id, "accion": "entregar", f"cant_{item.id}": "4"})
        ped.refresh_from_db()
        self.assertEqual(ped.estado, "entregado")
        self.assertEqual((saldo(self.t, self.m), self.m.stock_actual), (4, 16))
        self.assertTrue(Notificacion.objects.filter(persona=self.t, titulo__icontains="listo").exists())

    def test_supervisor_no_aprueba_pedidos_de_otro_equipo(self):
        from inventario.models import PedidoMaterial
        otro = persona("T2")
        ped = PedidoMaterial.objects.create(tecnico=otro)
        self.client.login(username="sup", password="x")
        self.assertEqual(self.client.post("/app/pedidos/", {"pedido": ped.id, "accion": "aprobar"}).status_code, 404)

    def test_partes_paradas_fifo(self):
        from inventario.stock_tecnico import consumir, entregar, partes_paradas
        entregar(self.t, self.m, 5, fecha=HOY - timedelta(days=90))
        entregar(self.t, self.m, 5, fecha=HOY - timedelta(days=5))
        consumir(self.t, None, self.m, 3, fecha=HOY - timedelta(days=2))  # consume primero lo más viejo
        f = partes_paradas([self.t], HOY)[0]
        self.assertEqual((f["cantidad"], f["dias"]), (2, 90))

    def test_faltante_para_ordenes(self):
        from inventario.models import RecetaMaterial
        from inventario.stock_tecnico import entregar, faltante_para_ordenes
        tipo = TipoTarea.objects.create(codigo="I", nombre="Inst")
        RecetaMaterial.objects.create(tipo_tarea=tipo, material=self.m, cantidad=1)
        for i in range(3):
            OrdenTrabajo.objects.create(numero=f"o{i}", tipo=tipo, tecnico=self.t, estado="asignada")
        entregar(self.t, self.m, 1)
        filas, n = faltante_para_ordenes(self.t)
        self.assertEqual((n, filas[0]["necesito"], filas[0]["falta"]), (3, 3, 2))

    def test_devolucion_vuelve_al_deposito(self):
        from inventario.stock_tecnico import devolver, entregar, saldo
        entregar(self.t, self.m, 5)
        devolver(self.t, self.m, 2)
        self.assertEqual((saldo(self.t, self.m), self.m.stock_actual), (3, 17))


class EncuestaSemanalTests(TestCase):
    def test_una_evaluacion_por_semana_y_cuenta_para_el_supervisor(self):
        from supervision.evaluacion import evaluar_supervisores
        from supervision.models import EncuestaSemanal
        sup = persona("S1", rol="supervisor")
        t = persona("T1", supervisor=sup, usuario=User.objects.create_user("tec", password="x"))
        self.client.login(username="tec", password="x")
        datos = {k: 4 for k in ("general", "trato", "organizacion", "apoyo", "ensenanza", "justicia")}
        self.client.post("/app/mi-supervisor/", datos)
        self.client.post("/app/mi-supervisor/", {**datos, "general": 1})  # segunda vez la misma semana: se ignora
        self.assertEqual(EncuestaSemanal.objects.filter(tecnico=t).count(), 1)
        ev = [e for e in evaluar_supervisores(HOY, 30) if e.supervisor == sup][0]
        self.assertEqual(ev.nota_semanal, 4)
        self.assertEqual(ev.score_imagen, 75)


class NotificacionesTests(TestCase):
    def test_orden_asignada_y_control_notifican_al_tecnico(self):
        from core.models import Notificacion
        sup = persona("S1", rol="supervisor")
        t = persona("T1", supervisor=sup, usuario=User.objects.create_user("tec", password="x"))
        tipo = TipoTarea.objects.create(codigo="R", nombre="Rep")
        OrdenTrabajo.objects.create(numero="1", tipo=tipo, tecnico=t, estado="asignada")
        InformeControl.objects.create(supervisor=sup, tecnico=t, puntaje=2, desvio_detectado=True)
        titulos = list(Notificacion.objects.filter(persona=t).values_list("titulo", flat=True))
        self.assertEqual(len(titulos), 2)
        self.client.login(username="tec", password="x")
        self.assertContains(self.client.get("/app/"), 'class="badge"')
        self.client.get("/app/notificaciones/")
        self.assertFalse(Notificacion.objects.filter(persona=t, leida=False).exists())


class IndicadoresProductividadTests(TestCase):
    """Cada indicador calcula lo que dice su definición (docs/METRICAS.md)."""

    def setUp(self):
        from datetime import datetime, time
        from personal.models import Asistencia
        self.sup = persona("S1", rol="supervisor")
        self.t = persona("T1", supervisor=self.sup)
        self.otro = persona("T2", supervisor=self.sup)
        self.tipo = TipoTarea.objects.create(codigo="I", nombre="Inst", minutos_estandar=60)
        self.dias = [HOY - timedelta(days=d) for d in range(1, 6)]
        tz = timezone.get_current_timezone()
        n = 0
        for d in self.dias:
            for t, cant in ((self.t, 4), (self.otro, 6)):
                Jornada.objects.create(fecha=d, tecnico=t)
                ent = timezone.make_aware(datetime.combine(d, time(8, 0)), tz)
                Asistencia.objects.create(persona=t, fecha=d, entrada=ent, salida=ent + timedelta(hours=8))
                for k in range(cant):
                    n += 1
                    OrdenTrabajo.objects.create(numero=f"o{n}", tipo=self.tipo, tecnico=t, estado="completada",
                                                fecha_programada=d, fecha_ejecucion=d, minutos_reales=60)

    def valores(self, tecnicos=None):
        from tablero.metricas import valores_tecnicos
        return valores_tecnicos(HOY - timedelta(days=29), HOY, tecnicos or [self.t, self.otro])

    def test_puntos_lineales_entre_minimo_y_meta(self):
        from core.models import Indicador
        mayor = Indicador(meta=70, minimo=40, mayor_es_mejor=True)
        menor = Indicador(meta=2, minimo=10, mayor_es_mejor=False)
        self.assertEqual((mayor.puntos(70), mayor.puntos(40), mayor.puntos(55), mayor.puntos(90)), (100, 0, 50, 100))
        self.assertEqual((menor.puntos(2), menor.puntos(10), menor.puntos(6), menor.puntos(0)), (100, 0, 50, 100))

    def test_eficiencia_de_la_jornada_en_horas_estandar(self):
        v = self.valores()
        self.assertAlmostEqual(v[self.t.id]["eficiencia_jornada"], 50)      # 4 h estándar de 8 h fichadas
        self.assertAlmostEqual(v[self.otro.id]["eficiencia_jornada"], 75)   # 6 de 8

    def test_primera_visita_descuenta_retrabajos(self):
        orig = OrdenTrabajo.objects.filter(tecnico=self.t).first()
        OrdenTrabajo.objects.create(numero="rt", tipo=self.tipo, tecnico=self.otro, es_retrabajo=True,
                                    orden_original=orig, estado="completada", fecha_programada=HOY, fecha_ejecucion=HOY)
        self.assertAlmostEqual(self.valores([self.t])[self.t.id]["primera_visita"], 95)  # 1 de 20

    def test_no_resueltas_evitables_ignora_causas_externas(self):
        d = self.dias[0]
        for k, motivo in enumerate(["cliente_ausente", "clima", "falta_material"]):
            OrdenTrabajo.objects.create(numero=f"f{k}", tipo=self.tipo, tecnico=self.t, estado="fallida",
                                        motivo_no_resuelto=motivo, fecha_programada=d, fecha_ejecucion=d)
        self.assertAlmostEqual(self.valores([self.t])[self.t.id]["no_resueltas_evitables"], 100 / 23)  # 1 de 23

    def test_cierre_en_sitio_por_distancia_gps(self):
        c = Cliente.objects.create(numero="c1", nombre="X", latitud=Decimal("-34.600000"), longitud=Decimal("-58.400000"))
        ots = list(OrdenTrabajo.objects.filter(tecnico=self.t)[:6])
        for k, o in enumerate(ots):
            o.cliente = c
            o.lat_cierre = Decimal("-34.600500") if k < 5 else Decimal("-34.620000")  # ~55 m vs ~2,2 km
            o.lng_cierre = Decimal("-58.400000")
            o.save()
        self.assertAlmostEqual(self.valores([self.t])[self.t.id]["cierre_en_sitio"], 500 / 6)

    def test_consumo_vs_estandar(self):
        from inventario.models import RecetaMaterial
        m = Material.objects.create(codigo="C", nombre="Cable", costo_unitario=Decimal("10"))
        RecetaMaterial.objects.create(tipo_tarea=self.tipo, material=m, cantidad=2)  # estándar $20 por orden
        for o in OrdenTrabajo.objects.filter(tecnico=self.t):
            Salida.objects.create(material=m, cantidad=3, costo_total=30, motivo="consumo", tecnico=self.t, orden=o)
        self.assertAlmostEqual(self.valores([self.t])[self.t.id]["consumo_vs_estandar"], 150)

    def test_indice_del_supervisor_y_tiempo_de_respuesta(self):
        from personal.models import Novedad
        from tablero.metricas import tableros_supervisores
        n = Novedad.objects.create(persona=self.t, tipo="enfermedad", estado="aprobada")
        n.resuelta = n.creada + timedelta(hours=6)
        n.save()
        tb = tableros_supervisores(supervisores=[self.sup])[0]
        self.assertAlmostEqual(tb.get("tiempo_respuesta").valor, 6, places=3)
        self.assertAlmostEqual(tb.get("eficiencia_equipo").valor, 62.5)
        self.assertEqual(tb.get("cobertura_control").valor, 0)  # trabajaron y nadie los controló
        self.assertIsNone(tb.get("respuesta_desvios").valor)    # menos de 3 desvíos: sin dato
        self.assertIsNotNone(tb.indice)

    def test_gerencia_puede_ajustar_metas(self):
        from django.core.management import call_command

        from core.models import Indicador
        call_command("configurar_grupos", stdout=io.StringIO())
        u = User.objects.create_user("ger", password="x", is_staff=True)
        u.groups.add(Group.objects.get(name="Gerencia"))
        self.client.login(username="ger", password="x")
        i = Indicador.objects.get(codigo="eficiencia_jornada")
        self.assertEqual(self.client.get(f"/admin/core/indicador/{i.id}/change/").status_code, 200)
        self.assertEqual(self.client.get("/tablero/productividad/").status_code, 200)
        self.assertEqual(self.client.get(f"/tablero/productividad/{self.t.id}/").status_code, 200)


class ExportacionExcelTests(TestCase):
    def test_panel_de_productividad_exporta_excel_con_graficos(self):
        from openpyxl import load_workbook
        sup = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="x"))
        persona("T1", supervisor=sup)
        u = User.objects.create_user("ger", password="x")
        u.groups.add(Group.objects.create(name="Gerencia"))
        self.client.login(username="ger", password="x")
        self.assertEqual(self.client.get("/tablero/productividad/general/").status_code, 200)
        r = self.client.get("/tablero/productividad/excel/?dias=30")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        wb = load_workbook(io.BytesIO(r.content))
        for hoja in ("Resumen", "Gráficos", "Indicadores empresa", "Técnicos", "Supervisores", "Por equipo"):
            self.assertIn(hoja, wb.sheetnames)
        self.assertGreaterEqual(len(wb["Gráficos"]._charts), 1)
        # el supervisor exporta sólo su equipo y no ve la hoja de supervisores
        self.client.login(username="sup", password="x")
        wb = load_workbook(io.BytesIO(self.client.get("/tablero/productividad/excel/").content))
        self.assertNotIn("Supervisores", wb.sheetnames)
        self.assertEqual(self.client.get("/tablero/productividad/general/").status_code, 403)



class LoginYCredencialesTests(TestCase):
    def setUp(self):
        self.ger = User.objects.create_user("ger", password="clave-segura-1")
        self.ger.groups.add(Group.objects.create(name="Gerencia"))
        self.sup = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="clave-segura-1"))
        self.t = persona("T7", supervisor=self.sup, dni="30111222",
                         usuario=User.objects.create_user("tecnico7", password="clave-segura-1"))

    def ingresar(self, usuario, clave="clave-segura-1"):
        return self.client.post("/login/", {"username": usuario, "password": clave})

    def test_ingreso_con_usuario_legajo_o_dni(self):
        for ident in ("tecnico7", "T7", "t7", "30111222"):
            self.client.logout()
            self.assertEqual(self.ingresar(ident).status_code, 302, ident)

    def test_bloqueo_tras_5_intentos_y_registro(self):
        from core.models import CuentaUsuario, RegistroAcceso
        for _ in range(5):
            self.ingresar("T7", "mala")
        r = self.ingresar("T7")  # aun con la clave correcta, está bloqueado
        self.assertContains(r, "bloqueado")
        self.assertTrue(CuentaUsuario.objects.get(usuario=self.t.usuario).bloqueada)
        self.assertEqual(RegistroAcceso.objects.filter(evento="fallido").count(), 5)
        self.assertTrue(RegistroAcceso.objects.filter(evento="bloqueo").exists())

    def test_blanqueo_obliga_a_cambiar_la_clave(self):
        from core.models import CuentaUsuario
        self.client.force_login(self.ger)
        self.client.post(f"/personal/usuarios/{self.t.id}/", {"accion": "blanquear"})
        clave = self.client.get(f"/personal/usuarios/{self.t.id}/").context["clave"]["clave"]
        self.assertEqual(len(clave), 14)
        self.client.logout()
        self.assertEqual(self.ingresar("T7", clave).status_code, 302)
        r = self.client.get("/app/")
        self.assertRedirects(r, "/cuenta/clave/?next=/app/", fetch_redirect_response=False)
        self.client.post("/cuenta/clave/", {"old_password": clave, "new_password1": "Otra-Clave-77",
                                            "new_password2": "Otra-Clave-77"})
        self.assertFalse(CuentaUsuario.objects.get(usuario=self.t.usuario).debe_cambiar_clave)
        self.assertEqual(self.client.get("/app/").status_code, 200)

    def test_crear_acceso_y_permisos_del_supervisor(self):
        nuevo = persona("T8", supervisor=self.sup)
        self.client.force_login(self.sup.usuario)
        self.client.post(f"/personal/usuarios/{nuevo.id}/", {"accion": "crear"})  # sólo gerencia crea accesos
        nuevo.refresh_from_db()
        self.assertIsNone(nuevo.usuario)
        r = self.client.post(f"/personal/usuarios/{self.t.id}/", {"accion": "blanquear"})  # su equipo: sí
        self.assertEqual(r.status_code, 302)
        ajeno = persona("T9")
        self.assertEqual(self.client.get(f"/personal/usuarios/{ajeno.id}/").status_code, 404)
        self.client.force_login(self.ger)
        self.client.post(f"/personal/usuarios/{nuevo.id}/", {"accion": "crear", "usuario": "t8"})
        nuevo.refresh_from_db()
        self.assertEqual(nuevo.usuario.username, "t8")
        self.assertTrue(nuevo.usuario.groups.filter(name="Técnicos").exists())

    def test_egreso_da_de_baja_el_acceso(self):
        self.t.activo, self.t.fecha_egreso = False, HOY
        self.t.save()
        self.t.usuario.refresh_from_db()
        self.assertFalse(self.t.usuario.is_active)
        self.assertContains(self.ingresar("T7"), "dado de baja")

    def test_privacidad_se_pide_a_tecnicos(self):
        from core.models import CuentaUsuario
        CuentaUsuario.objects.filter(usuario=self.t.usuario).update(acepto_privacidad=None)
        self.client.force_login(self.t.usuario)
        self.assertRedirects(self.client.get("/app/"), "/cuenta/privacidad/?next=/app/", fetch_redirect_response=False)
        self.client.post("/cuenta/privacidad/", {"acepto": "1", "next": "/app/"})
        self.assertEqual(self.client.get("/app/").status_code, 200)

    def test_rol_deposito_solo_inventario(self):
        from django.core.management import call_command
        call_command("configurar_grupos", stdout=io.StringIO())
        dep = User.objects.create_user("dep", password="x", is_staff=True)
        dep.groups.add(Group.objects.get(name="Depósito"))
        self.client.force_login(dep)
        self.assertRedirects(self.client.get("/"), "/tablero/pedidos/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/tablero/pedidos/").status_code, 200)
        self.assertEqual(self.client.get("/tablero/stock/").status_code, 200)
        self.assertEqual(self.client.get("/tablero/finanzas/").status_code, 403)
        self.assertEqual(self.client.get("/tablero/tecnicos/").status_code, 403)


class PartesAdeudadasTests(TestCase):
    def setUp(self):
        self.sup = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="x"))
        self.t = persona("T1", supervisor=self.sup, usuario=User.objects.create_user("tec", password="x"))
        self.tipo = TipoTarea.objects.create(codigo="RET", nombre="Retiro")
        self.ger = User.objects.create_user("ger", password="x")
        self.ger.groups.add(Group.objects.create(name="Gerencia"))

    def test_cerrar_retiro_registra_equipos_y_genera_deuda(self):
        from inventario.deudas import deudas
        from inventario.models import EquipoRetirado
        ot = OrdenTrabajo.objects.create(numero="1", tipo=self.tipo, tecnico=self.t, estado="asignada")
        self.client.login(username="tec", password="x")
        self.client.post(f"/app/orden/{ot.id}/", {"resultado": "completada", "series_retiradas": "ABC1, ABC2"})
        self.assertEqual(EquipoRetirado.objects.filter(tecnico=self.t, estado="en_tecnico").count(), 2)
        EquipoRetirado.objects.update(fecha_retiro=HOY - timedelta(days=8))
        d = deudas([self.t])
        self.assertEqual((len(d), d[0].dias, d[0].estado), (2, 8, "critico"))  # más de 5 días: rojo
        self.assertContains(self.client.get("/app/"), "por regularizar")

    def test_aviso_diario_y_regularizacion_en_deposito(self):
        from core.models import Notificacion
        from inventario.deudas import deudas, notificar_deudas
        from inventario.models import EquipoRetirado
        e = EquipoRetirado.objects.create(tecnico=self.t, numero_serie="X1", fecha_retiro=HOY - timedelta(days=3))
        notificar_deudas(log=lambda *_: None)
        self.assertTrue(Notificacion.objects.filter(persona=self.t, titulo__icontains="regularizar").exists())
        self.assertTrue(Notificacion.objects.filter(persona=self.sup, titulo__icontains="tu equipo").exists())
        self.assertEqual(deudas([self.t])[0].estado, "aviso")  # 3 días: amarillo
        self.client.force_login(self.ger)
        self.client.post("/tablero/partes-adeudadas/", {"accion": "recibir", "equipo": e.id})
        e.refresh_from_db()
        self.assertEqual(e.estado, "devuelto")
        self.assertEqual(deudas([self.t]), [])

    def test_uso_sin_cargo_cuenta_desde_que_quedo_negativo(self):
        from inventario.deudas import deudas
        from inventario.stock_tecnico import consumir, entregar
        m = Material.objects.create(codigo="R", nombre="Roseta", costo_unitario=Decimal("1"))
        LoteIngreso.objects.create(material=m, cantidad=100)
        entregar(self.t, m, 2, fecha=HOY - timedelta(days=10))
        consumir(self.t, None, m, 5, fecha=HOY - timedelta(days=4))
        d = deudas([self.t])[0]
        self.assertEqual((d.tipo, d.cantidad, d.dias), ("uso_sin_cargo", 3, 4))

    def test_supervisor_ve_su_equipo_pero_no_regulariza(self):
        from inventario.models import EquipoRetirado
        ajeno = persona("T2")
        EquipoRetirado.objects.create(tecnico=ajeno, numero_serie="Z9")
        e = EquipoRetirado.objects.create(tecnico=self.t, numero_serie="Y1")
        self.client.force_login(self.sup.usuario)
        r = self.client.get("/tablero/partes-adeudadas/")
        self.assertContains(r, "Y1")
        self.assertNotContains(r, "Z9")
        self.client.post("/tablero/partes-adeudadas/", {"accion": "recibir", "equipo": e.id})
        e.refresh_from_db()
        self.assertEqual(e.estado, "en_tecnico")


class MetasYSemaforosTests(TestCase):
    def setUp(self):
        self.ger = User.objects.create_user("ger", password="x")
        self.ger.groups.add(Group.objects.create(name="Gerencia"))
        self.sup = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="x"))
        self.otro_sup = persona("S2", rol="supervisor")
        self.t1 = persona("T1", supervisor=self.sup)
        self.t2 = persona("T2", supervisor=self.otro_sup)

    def test_regla_verde_amarillo_rojo(self):
        from core.models import Indicador
        i = Indicador(meta=90, minimo=70, mayor_es_mejor=True, unidad="%")
        self.assertEqual([i.estado(v) for v in (95, 90, 80, 70, 60, None)], ["ok", "ok", "aviso", "critico", "critico", "info"])
        j = Indicador(meta=0, minimo=3, mayor_es_mejor=False, unidad="personas")
        self.assertEqual([j.estado(v) for v in (0, 2, 3)], ["ok", "aviso", "critico"])

    def test_supervisor_fija_metas_de_su_equipo_y_queda_registro(self):
        from core.metas import indicadores_efectivos
        from core.models import CambioMeta, Indicador
        i = Indicador.objects.get(codigo="primera_visita")
        self.client.force_login(self.sup.usuario)
        self.client.post("/tablero/metas/", {f"meta_{i.id}": "98", f"lim_{i.id}": "90"})
        propia = {x.codigo: x for x in indicadores_efectivos("tecnico", self.sup)}["primera_visita"]
        general = {x.codigo: x for x in indicadores_efectivos("tecnico", self.otro_sup)}["primera_visita"]
        self.assertEqual((propia.meta, propia.minimo), (98, 90))
        self.assertEqual((general.meta, general.minimo), (95, 80))  # el otro equipo usa la general
        self.assertTrue(CambioMeta.objects.filter(indicador=i, usuario=self.sup.usuario).exists())

    def test_supervisor_no_puede_tocar_las_metas_que_lo_evaluan(self):
        from core.models import Indicador, MetaEquipo
        i = Indicador.objects.get(codigo="clima_equipo")
        self.client.force_login(self.sup.usuario)
        self.client.post("/tablero/metas/", {f"meta_{i.id}": "10", f"lim_{i.id}": "1"})
        self.assertFalse(MetaEquipo.objects.filter(indicador=i).exists())

    def test_meta_incoherente_se_rechaza(self):
        from core.models import Indicador
        i = Indicador.objects.get(codigo="presentismo_hoy")
        self.client.force_login(self.ger)
        r = self.client.post("/tablero/metas/", {f"meta_{i.id}": "80", f"lim_{i.id}": "90"}, follow=True)
        self.assertContains(r, "debe ser mayor")
        i.refresh_from_db()
        self.assertEqual(i.meta, 95)

    def test_resumen_muestra_semaforos(self):
        self.client.force_login(self.ger)
        r = self.client.get("/tablero/")
        self.assertContains(r, "card kpi mando")
        self.client.force_login(self.sup.usuario)
        self.assertEqual(self.client.get("/tablero/").status_code, 200)


class DebriefTests(TestCase):
    def setUp(self):
        self.sup = persona("S1", rol="supervisor")
        self.t = persona("T1", supervisor=self.sup, usuario=User.objects.create_user("tec", password="x"))
        self.otro = persona("T2", supervisor=self.sup)
        self.cli = Cliente.objects.create(numero="C1", nombre="Ana", telefono="1155550000", email="ana@x.com",
                                          direccion="Calle 1")
        self.tipo = TipoTarea.objects.create(codigo="I", nombre="Instalación")
        self.ot = OrdenTrabajo.objects.create(numero="OT1", tipo=self.tipo, tecnico=self.t, cliente=self.cli,
                                              estado="asignada", observaciones="Timbre 2B")
        self.client.login(username="tec", password="x")

    def test_ficha_muestra_tarea_contacto_lom_y_notas_de_visitas_anteriores(self):
        from operaciones.models import NotaOrden
        anterior = OrdenTrabajo.objects.create(numero="OT0", tipo=self.tipo, tecnico=self.otro, cliente=self.cli,
                                               estado="completada", fecha_ejecucion=HOY - timedelta(days=20))
        NotaOrden.objects.create(orden=anterior, autor=self.otro, tipo="cierre", texto="Se cambió el conector")
        r = self.client.get(f"/app/orden/{self.ot.id}/")
        for texto in ("Instalación", "1155550000", "ana@x.com", "Calle 1", "LOM", "Notas (2)"):
            self.assertContains(r, texto)
        r = self.client.get(f"/app/orden/{self.ot.id}/notas/")
        self.assertContains(r, "Se cambió el conector")
        self.assertContains(r, "visita anterior")
        self.client.post(f"/app/orden/{self.ot.id}/notas/", {"texto": "Llegué, el cliente no está"})
        self.assertTrue(NotaOrden.objects.filter(orden=self.ot, texto__icontains="no está").exists())

    def test_debrief_completo(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from finanzas.models import Egreso
        from operaciones.models import GastoOrden, NotaOrden
        r = self.client.post(f"/app/orden/{self.ot.id}/cierre/", {
            "resultado": "completada", "nota": "Instalé ONT y probé velocidad", "minutos_viaje": "25",
            "minutos_reales": "70", "minutos_retorno": "20", "conforme_nombre": "Ana", "conforme_apellido": "Gómez",
            "conforme_dni": "30111222", "gasto_desc_1": "Tarugos", "gasto_monto_1": "3500",
            "gasto_comp_1": SimpleUploadedFile("t.jpg", b"img", content_type="image/jpeg")})
        self.assertEqual(r.status_code, 302)
        self.ot.refresh_from_db()
        self.assertEqual((self.ot.estado, self.ot.minutos_viaje, self.ot.minutos_reales, self.ot.minutos_retorno),
                         ("completada", 25, 70, 20))
        self.assertEqual((self.ot.conforme_nombre, self.ot.conforme_apellido), ("Ana", "Gómez"))
        self.assertEqual(NotaOrden.objects.get(orden=self.ot, tipo="cierre").texto, "Instalé ONT y probé velocidad")
        g = GastoOrden.objects.get(orden=self.ot)
        self.assertEqual(Egreso.objects.get(origen=f"gasto_orden:{g.id}").monto, Decimal("3500"))
        g.comprobante.delete()

    def test_gasto_sin_monto_se_rechaza(self):
        r = self.client.post(f"/app/orden/{self.ot.id}/cierre/", {"resultado": "completada", "gasto_desc_1": "Cinta"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Completá qué compraste")

    def test_lom_pide_partes_para_la_llamada_descontando_lo_que_tiene(self):
        from inventario.models import PedidoMaterial, RecetaMaterial
        from inventario.stock_tecnico import entregar
        m = Material.objects.create(codigo="ONT", nombre="Módem")
        LoteIngreso.objects.create(material=m, cantidad=10)
        RecetaMaterial.objects.create(tipo_tarea=self.tipo, material=m, cantidad=3)
        entregar(self.t, m, 1)
        r = self.client.get(f"/app/stock/pedir/?orden={self.ot.id}")
        self.assertEqual(r.context["form"].initial["cantidad_1"], 2)  # necesita 3, tiene 1
        self.client.post("/app/stock/pedir/", {"orden": self.ot.id, "material_1": m.id, "cantidad_1": "2"})
        self.assertEqual(PedidoMaterial.objects.get().orden, self.ot)


class ContabilidadTests(TestCase):
    def setUp(self):
        from django.core.management import call_command
        call_command("configurar_grupos", stdout=io.StringIO())
        self.cont = User.objects.create_user("cont", password="x", is_staff=True)
        self.cont.groups.add(Group.objects.get(name="Contabilidad"))
        self.client.force_login(self.cont)

    def test_contabilidad_entra_a_finanzas_y_no_al_resto(self):
        self.assertRedirects(self.client.get("/"), "/finanzas/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/finanzas/").status_code, 200)
        self.assertEqual(self.client.get("/tablero/finanzas/").status_code, 200)
        self.assertEqual(self.client.get("/tablero/tecnicos/").status_code, 403)
        self.assertEqual(self.client.get("/personal/legajos/").status_code, 403)

    def test_cargar_factura_con_proveedor_nuevo_pagarla_y_eliminarla(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from finanzas.models import CategoriaEgreso, Egreso, Proveedor
        cat = CategoriaEgreso.objects.create(nombre="Servicios", codigo="servicios")
        r = self.client.post("/finanzas/comprobantes/nuevo/", {
            "tipo_comprobante": "factura", "numero_comprobante": "A-0001-00000012", "fecha": HOY.isoformat(),
            "proveedor_nombre": "Edesur", "cuit": "30-65511620-2", "categoria": cat.id, "monto": "125.000,50",
            "vencimiento": (HOY + timedelta(days=10)).isoformat(),
            "comprobante": SimpleUploadedFile("f.pdf", b"%PDF-1.1", "application/pdf")})
        self.assertEqual(r.status_code, 302)
        e = Egreso.objects.get()
        self.assertEqual((e.monto, e.proveedor_ref.cuit, e.pagado, e.cargado_por), (Decimal("125000.50"), "30-65511620-2",
                                                                                   False, self.cont))
        self.assertTrue(e.comprobante.name.endswith(".pdf"))
        self.assertEqual(Proveedor.objects.get().categoria, cat)  # queda como categoría habitual
        # segunda factura del mismo proveedor sin categoría → se usa la habitual; la nota de crédito resta
        self.client.post("/finanzas/comprobantes/nuevo/", {"tipo_comprobante": "nota_credito", "fecha": HOY.isoformat(),
                                                          "proveedor_nombre": "Edesur", "monto": "5000", "pagado": "on"})
        nc = Egreso.objects.get(tipo_comprobante="nota_credito")
        self.assertEqual((nc.categoria, nc.monto, nc.fecha_pago), (cat, Decimal("-5000"), HOY))
        self.assertEqual(Proveedor.objects.count(), 1)
        self.assertContains(self.client.get("/finanzas/comprobantes/?ver=a_pagar"), "A-0001-00000012")
        self.client.post("/finanzas/comprobantes/", {"accion": "pagar", "egreso": e.id, "medio_pago": "cheque"})
        e.refresh_from_db()
        self.assertEqual((e.pagado, e.fecha_pago, e.medio_pago), (True, HOY, "cheque"))
        self.client.post(f"/finanzas/comprobantes/{e.id}/", {"accion": "eliminar"})
        self.assertFalse(Egreso.objects.filter(pk=e.id).exists())

    def test_numeros_como_se_escriben_aca(self):
        from finanzas.views import _dec
        self.assertEqual([_dec(x) for x in ("125.000,50", "10.000", "1135000.00", "27,5", "", "abc")],
                         [Decimal("125000.50"), Decimal("10000"), Decimal("1135000.00"), Decimal("27.5"), None, False])

    def test_aprobar_y_rechazar_reintegro(self):
        from finanzas.models import Egreso
        from operaciones.models import GastoOrden
        t = persona("T1")
        ot = OrdenTrabajo.objects.create(numero="1", tipo=TipoTarea.objects.create(codigo="R", nombre="R"), tecnico=t)
        g = GastoOrden.objects.create(orden=ot, tecnico=t, descripcion="Cinta", monto=2000)
        self.assertTrue(Egreso.objects.filter(origen=f"gasto_orden:{g.id}").exists())
        self.client.post("/finanzas/reintegros/", {"gasto": g.id, "accion": "rechazar", "motivo": "sin factura"})
        g.refresh_from_db()
        self.assertEqual(g.estado, "rechazado")
        self.assertFalse(Egreso.objects.filter(origen=f"gasto_orden:{g.id}").exists())

    def test_presupuesto_con_semaforo(self):
        from finanzas.models import CategoriaEgreso, Egreso, Presupuesto
        from finanzas.views import presupuesto_vs_real
        mes = HOY.replace(day=1)
        c = CategoriaEgreso.objects.create(nombre="Flota", codigo="flota", tolerancia_presupuesto=10)
        Presupuesto.objects.create(categoria=c, mes=mes, monto=1000)
        for monto, esperado in ((900, "ok"), (150, "aviso"), (100, "critico")):  # 900 → 1050 → 1150
            Egreso.objects.create(categoria=c, fecha=mes, monto=monto, descripcion="x")
            self.assertEqual(presupuesto_vs_real(mes)[0]["estado"], esperado)
        self.client.post(f"/finanzas/presupuesto/?mes={mes:%Y-%m}", {f"p_{c.id}_{mes:%Y%m}": "2.000,50"})
        self.assertEqual(Presupuesto.objects.get(categoria=c, mes=mes).monto, Decimal("2000.50"))


class IndicadoresEditablesTests(TestCase):
    def setUp(self):
        self.ger = User.objects.create_user("ger", password="x")
        self.ger.groups.add(Group.objects.create(name="Gerencia"))
        self.sup = persona("S1", rol="supervisor", usuario=User.objects.create_user("sup", password="x"))
        self.t = persona("T1", supervisor=self.sup)
        self.ajeno = persona("T2")
        self.client.force_login(self.ger)

    def test_crear_indicador_manual_cargar_valores_y_que_cuente(self):
        from core.models import Indicador
        from tablero.metricas import tableros_tecnicos
        r = self.client.post("/tablero/metas/indicador/nuevo/", {
            "rol": "tecnico", "nombre": "Reclamos de clientes", "descripcion": "Reclamos recibidos en el mes",
            "unidad": "reclamos", "meta": "2", "minimo": "6", "peso": "10", "activo": "on"})
        i = Indicador.objects.get(nombre="Reclamos de clientes")
        self.assertEqual((i.tipo, i.mayor_es_mejor), ("manual", False))
        self.assertRedirects(r, f"/tablero/metas/indicador/{i.id}/valores/", fetch_redirect_response=False)
        self.client.post(f"/tablero/metas/indicador/{i.id}/valores/", {"mes": f"{HOY:%Y-%m}", f"v_{self.t.id}": "4"})
        m = tableros_tecnicos(tecnicos=[self.t])[0].get(i.codigo)
        self.assertEqual((m.valor, m.estado), (4, "aviso"))

    def test_supervisor_carga_valores_solo_de_su_equipo(self):
        from core.models import Indicador, ValorIndicador
        i = Indicador.objects.create(codigo="manual-x", rol="tecnico", tipo="manual", nombre="X", descripcion="x",
                                     meta=10, minimo=5)
        self.client.force_login(self.sup.usuario)
        self.client.post(f"/tablero/metas/indicador/{i.id}/valores/", {f"v_{self.t.id}": "8", f"v_{self.ajeno.id}": "1"})
        self.assertEqual(list(ValorIndicador.objects.values_list("persona_id", flat=True)), [self.t.id])
        self.assertEqual(self.client.get("/tablero/metas/indicador/nuevo/").status_code, 403)

    def test_editar_desactivar_y_no_eliminar_los_del_sistema(self):
        from core.models import Indicador
        i = Indicador.objects.get(codigo="primera_visita")
        self.client.post(f"/tablero/metas/indicador/{i.id}/", {
            "nombre": "Sin retrabajos", "descripcion": i.descripcion, "unidad": "%", "mayor_es_mejor": "on",
            "meta": "96", "minimo": "85", "peso": "15"})  # sin 'activo' → se desactiva
        i.refresh_from_db()
        self.assertEqual((i.nombre, i.peso, i.activo), ("Sin retrabajos", 15, False))
        self.client.post(f"/tablero/metas/indicador/{i.id}/", {"accion": "eliminar"})
        self.assertTrue(Indicador.objects.filter(pk=i.pk).exists())
        r = self.client.post(f"/tablero/metas/indicador/{i.id}/", {
            "nombre": "x", "descripcion": "x", "unidad": "%", "mayor_es_mejor": "on", "meta": "50", "minimo": "80",
            "peso": "5", "activo": "on"})
        self.assertContains(r, "lado")  # meta incoherente con el límite


class SueldosTests(TestCase):
    """Pre-liquidación: horas extra 50/100, presentismo, multas informativas (art. 131 LCT), suspensiones."""

    def setUp(self):
        self.p = Parametros.actual()
        self.p.horas_mensuales, self.p.recargo_extra_50, self.p.recargo_extra_100 = 200, 50, 100
        self.p.adicional_presentismo, self.p.presentismo_tardanzas_max, self.p.cargas_sociales = Decimal("10"), 2, Decimal("25")
        self.p.save()
        self.t = persona("T1", sueldo_basico=Decimal("1000000"))
        self.mes = (HOY.replace(day=1) - timedelta(days=1)).replace(day=1)  # mes anterior completo

    def _fichar(self, dia, horas_extra):
        from personal.models import Asistencia
        entrada = timezone.make_aware(timezone.datetime.combine(dia, timezone.datetime.min.time())) + timedelta(hours=8)
        a = Asistencia.objects.create(persona=self.t, fecha=dia, entrada=entrada, salida=entrada + timedelta(hours=9))
        Asistencia.objects.filter(pk=a.pk).update(horas_extra=horas_extra)

    def _calcular(self, tardanzas=0, injustificadas=0):
        from types import SimpleNamespace
        from unittest import mock

        from finanzas.sueldos import calcular
        asis = SimpleNamespace(presentes=20, injustificadas=injustificadas, tardanzas=tardanzas)
        with mock.patch("finanzas.sueldos.resumen", return_value=[asis]):
            return calcular(self.t, self.mes)

    def test_horas_extra_presentismo_multa_y_suspension(self):
        dias = [self.mes + timedelta(days=i) for i in range(14)]
        habil = next(d for d in dias if d.weekday() < 5)
        domingo = next(d for d in dias if d.weekday() == 6)
        self._fichar(habil, Decimal("4"))
        self._fichar(domingo, Decimal("2"))
        AccionCorrectiva.objects.create(tecnico=self.t, tipo="multa", monto=40000, fecha=habil, descripcion="x")
        AccionCorrectiva.objects.create(tecnico=self.t, tipo="suspension", dias_suspension=2, fecha=habil, descripcion="x")
        liq = self._calcular(tardanzas=1)
        # valor hora 5.000 → 4 h al 50 % = 30.000 + 2 h al 100 % = 20.000
        self.assertEqual((liq.horas_extra_50, liq.horas_extra_100, liq.monto_horas_extra), (4, 2, Decimal("50000.00")))
        self.assertEqual(liq.presentismo, Decimal("100000.00"))
        self.assertEqual(liq.multas, 40000)  # se informa…
        self.assertEqual(liq.descuento_dias, Decimal("66666.67"))  # …pero sólo se descuentan los días de suspensión
        self.assertEqual(liq.bruto, Decimal("1083333.33"))
        self.assertEqual(liq.costo_total, Decimal("1354166.66"))

    def test_pierde_presentismo_por_llegadas_tarde_o_faltas(self):
        self.assertEqual(self._calcular(tardanzas=3).presentismo, 0)
        liq = self._calcular(injustificadas=1)
        self.assertEqual((liq.presentismo, liq.dias_descuento, liq.descuento_dias), (0, 1, Decimal("33333.33")))

    def test_flujo_generar_ajustar_aprobar_y_gasto(self):
        from django.core.management import call_command

        from finanzas.models import Egreso, Liquidacion
        from finanzas.views import bloques_del_mes
        call_command("configurar_grupos", stdout=io.StringIO())
        cont = User.objects.create_user("cont", password="x")
        cont.groups.add(Group.objects.get(name="Contabilidad"))
        self.client.force_login(cont)
        url = f"/finanzas/sueldos/?mes={self.mes:%Y-%m}"
        self.client.post(url, {"accion": "generar", "mes": f"{self.mes:%Y-%m}"})
        liq = Liquidacion.objects.get(persona=self.t)
        self.assertEqual(liq.estado, "borrador")
        self.assertFalse(Egreso.objects.filter(origen=f"liquidacion:{liq.id}").exists())  # el borrador no es gasto
        bruto = liq.bruto
        self.client.post(url, {"accion": "fila", "liq": liq.id, "mes": f"{self.mes:%Y-%m}", "otros_adicionales": "10.000",
                               "otros_descuentos": "", "estado": "borrador"})
        liq.refresh_from_db()
        self.assertEqual(liq.bruto, bruto + 10000)
        self.client.post(url, {"accion": "aprobar_todas", "mes": f"{self.mes:%Y-%m}"})
        liq.refresh_from_db()
        e = Egreso.objects.get(origen=f"liquidacion:{liq.id}")
        self.assertEqual((liq.estado, e.monto, e.categoria.codigo), ("aprobada", liq.costo_total, "sueldos"))
        b = bloques_del_mes(self.mes)
        self.assertEqual(b["sueldos"] + b["horas_extra"], liq.costo_total)
        r = self.client.get(url + "&formato=excel")
        self.assertEqual(r["Content-Type"].split(";")[0], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        # los supervisores no ven sueldos
        sup = User.objects.create_user("sup", password="x")
        persona("S1", rol="supervisor", usuario=sup)
        self.client.force_login(sup)
        self.assertEqual(self.client.get(url).status_code, 403)
