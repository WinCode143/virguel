"""Importación masiva desde Excel (.xlsx) o CSV para la carga inicial.

Cada importador define sus columnas. Todo el archivo se procesa en una sola
transacción: si alguna fila tiene error, no se guarda nada y se informa
fila por fila qué corregir. Las filas existentes (misma clave) se actualizan.
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.models import Group, User
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render

from core.models import Cliente, Persona, Zona
from core.roles import GERENCIA, requiere_rol
from flota.models import Vehiculo
from inventario.models import DemandaComercial, LoteIngreso, Material
from operaciones.models import OrdenTrabajo, TipoTarea


class ErrorFila(Exception):
    pass


# ---------------------------------------------------------------- conversores
def texto(v, obligatorio=True):
    v = "" if v is None else str(v).strip()
    if obligatorio and not v:
        raise ErrorFila("valor obligatorio vacío")
    return v


def entero(v, obligatorio=True):
    v = texto(v, obligatorio)
    if not v:
        return None
    try:
        return int(float(v.replace(",", ".")))
    except ValueError:
        raise ErrorFila(f"'{v}' no es un número entero")


def decimal_(v, obligatorio=True):
    v = texto(v, obligatorio)
    if not v:
        return None
    v = v.replace("$", "").replace(" ", "")
    if "," in v:  # formato argentino 1.234,56
        v = v.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", v):  # 1.500 = mil quinientos
        v = v.replace(".", "")
    try:
        return Decimal(v)
    except InvalidOperation:
        raise ErrorFila(f"'{v}' no es un número")


def fecha(v, obligatorio=True):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    v = texto(v, obligatorio)
    if not v:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            pass
    raise ErrorFila(f"fecha '{v}' inválida (usar dd/mm/aaaa)")


def si_no(v):
    return texto(v, False).lower() in ("si", "sí", "s", "1", "true", "x", "verdadero")


def zona(v):
    v = texto(v, False)
    return Zona.objects.get_or_create(nombre=v)[0] if v else None


def persona_por_legajo(v, rol=None, obligatorio=True):
    v = texto(v, obligatorio)
    if not v:
        return None
    try:
        return Persona.objects.get(legajo=v, **({"rol": rol} if rol else {}))
    except Persona.DoesNotExist:
        raise ErrorFila(f"no existe {'el ' + rol if rol else 'la persona'} con legajo '{v}'")


def elegir(v, opciones, nombre):
    v = texto(v).lower()
    for valor, etiqueta in opciones:
        if v in (valor.lower(), etiqueta.lower()):
            return valor
    raise ErrorFila(f"{nombre} '{v}' inválido (opciones: {', '.join(o[0] for o in opciones)})")


# ---------------------------------------------------------------- importadores
def imp_personas(f):
    rol = elegir(f["rol"], Persona.Rol.choices, "rol")
    p, creado = Persona.objects.update_or_create(legajo=texto(f["legajo"]), defaults={
        "nombre": texto(f["nombre"]), "apellido": texto(f["apellido"]), "dni": texto(f.get("dni"), False),
        "rol": rol, "zona": zona(f.get("zona")), "telefono": texto(f.get("telefono"), False),
        "email": texto(f.get("email"), False), "fecha_ingreso": fecha(f.get("fecha_ingreso"), False) or date.today(),
        "supervisor": persona_por_legajo(f.get("legajo_supervisor"), "supervisor", False)})
    usuario = texto(f.get("usuario"), False)
    if usuario and not p.usuario_id:
        u, nuevo = User.objects.get_or_create(username=usuario, defaults={"first_name": p.nombre, "last_name": p.apellido})
        if nuevo:
            u.set_password(texto(f.get("clave_inicial"), False) or p.dni or p.legajo)
            u.save()
        grupo = {"tecnico": "Técnicos", "supervisor": "Supervisores", "gerencia": "Gerencia",
                 "administrativo": "Administración"}[rol]
        u.groups.add(Group.objects.get_or_create(name=grupo)[0])
        p.usuario = u
        p.save(update_fields=["usuario"])
    return creado


def imp_vehiculos(f):
    _, creado = Vehiculo.objects.update_or_create(patente=texto(f["patente"]).upper().replace(" ", ""), defaults={
        "marca": texto(f["marca"]), "modelo": texto(f["modelo"]), "anio": entero(f["anio"]),
        "km_actual": entero(f.get("km_actual"), False) or 0,
        "asignado_a": persona_por_legajo(f.get("legajo_asignado"), obligatorio=False),
        "vencimiento_vtv": fecha(f.get("vencimiento_vtv"), False),
        "vencimiento_seguro": fecha(f.get("vencimiento_seguro"), False)})
    return creado


def imp_materiales(f):
    _, creado = Material.objects.update_or_create(codigo=texto(f["codigo"]), defaults={
        "nombre": texto(f["nombre"]),
        "categoria": elegir(f.get("categoria") or "material", Material.Categoria.choices, "categoría"),
        "unidad": texto(f.get("unidad"), False) or "u", "costo_unitario": decimal_(f["costo_unitario"]),
        "stock_minimo": decimal_(f.get("stock_minimo"), False) or 0,
        "es_decodificador": si_no(f.get("es_decodificador"))})
    return creado


def imp_stock(f):
    try:
        m = Material.objects.get(codigo=texto(f["codigo_material"]))
    except Material.DoesNotExist:
        raise ErrorFila(f"no existe el material '{f['codigo_material']}' (importar materiales primero)")
    LoteIngreso.objects.create(material=m, fecha=fecha(f["fecha_ingreso"]), cantidad=decimal_(f["cantidad"]),
                               costo_unitario=decimal_(f.get("costo_unitario"), False) or m.costo_unitario,
                               proveedor=texto(f.get("proveedor"), False), remito=texto(f.get("remito"), False))
    return True


def tipo_tarea(v):
    try:
        return TipoTarea.objects.get(codigo=texto(v))
    except TipoTarea.DoesNotExist:
        raise ErrorFila(f"no existe el tipo de tarea '{v}'")


def imp_demanda(f):
    DemandaComercial.objects.create(fecha=fecha(f["fecha"]), zona=zona(f.get("zona")),
                                    tipo_tarea=tipo_tarea(f["codigo_tipo_tarea"]),
                                    cantidad_clientes=entero(f["cantidad_clientes"]),
                                    observaciones=texto(f.get("observaciones"), False))
    return True


def imp_clientes(f):
    _, creado = Cliente.objects.update_or_create(numero=texto(f["numero"]), defaults={
        "nombre": texto(f["nombre"]), "direccion": texto(f.get("direccion"), False), "zona": zona(f.get("zona")),
        "tipo": elegir(f.get("tipo") or "residencial", Cliente.Tipo.choices, "tipo"),
        "cantidad_televisores": entero(f.get("cantidad_televisores"), False) or 1})
    return creado


def imp_ordenes(f):
    cli = None
    if texto(f.get("numero_cliente"), False):
        cli = Cliente.objects.filter(numero=texto(f["numero_cliente"])).first()
        if cli is None:
            raise ErrorFila(f"no existe el cliente '{f['numero_cliente']}'")
    tec = persona_por_legajo(f.get("legajo_tecnico"), "tecnico", False)
    _, creado = OrdenTrabajo.objects.update_or_create(numero=texto(f["numero"]), defaults={
        "tipo": tipo_tarea(f["codigo_tipo_tarea"]), "cliente": cli, "tecnico": tec,
        "zona": zona(f.get("zona")) or (cli.zona if cli else None),
        "fecha_programada": fecha(f["fecha_programada"]), "estado": "asignada" if tec else "pendiente",
        "observaciones": texto(f.get("observaciones"), False)})
    return creado


IMPORTADORES = {
    "personas": ("Personas (técnicos, supervisores, administración)", imp_personas,
                 ["legajo", "nombre", "apellido", "rol", "legajo_supervisor", "zona", "dni", "telefono", "email",
                  "fecha_ingreso", "usuario", "clave_inicial"],
                 ["T001", "Juan", "Pérez", "tecnico", "S001", "Norte", "30111222", "1155550000", "",
                  "01/03/2024", "jperez", ""],
                 "Importar primero los supervisores y después los técnicos. Si se indica 'usuario', se crea el acceso "
                 "(clave inicial: la indicada, o el DNI, o el legajo)."),
    "vehiculos": ("Vehículos", imp_vehiculos,
                  ["patente", "marca", "modelo", "anio", "km_actual", "legajo_asignado", "vencimiento_vtv",
                   "vencimiento_seguro"],
                  ["AB123CD", "Renault", "Kangoo", "2021", "85000", "T001", "15/11/2026", "01/02/2027"], ""),
    "materiales": ("Materiales y equipos", imp_materiales,
                   ["codigo", "nombre", "categoria", "unidad", "costo_unitario", "stock_minimo", "es_decodificador"],
                   ["DECO-HD", "Decodificador TV HD", "equipo", "u", "39000", "40", "si"],
                   "categoría: material, equipo o repuesto."),
    "stock": ("Stock actual (lotes con fecha de ingreso)", imp_stock,
              ["codigo_material", "fecha_ingreso", "cantidad", "costo_unitario", "proveedor", "remito"],
              ["DECO-HD", "10/08/2026", "120", "39000", "Proveedor SA", "R-0001"],
              "Una fila por ingreso: la fecha real de ingreso es la que permite controlar los 60 días."),
    "clientes": ("Clientes", imp_clientes,
                 ["numero", "nombre", "direccion", "zona", "tipo", "cantidad_televisores"],
                 ["C000123", "María Gómez", "Calle 12 N° 345", "Centro", "moderno", "2"],
                 "tipo: residencial, moderno o comercial."),
    "demanda": ("Demanda comercial prevista", imp_demanda,
                ["fecha", "codigo_tipo_tarea", "cantidad_clientes", "zona", "observaciones"],
                ["20/10/2026", "INST-FO", "45", "Norte", "Campaña fibra"], ""),
    "ordenes": ("Órdenes de trabajo", imp_ordenes,
                ["numero", "codigo_tipo_tarea", "fecha_programada", "legajo_tecnico", "numero_cliente", "zona",
                 "observaciones"],
                ["OT0001", "INST-FO", "08/10/2026", "T001", "C000123", "", "Timbre 2B"], ""),
}


OBLIGATORIAS = {
    "personas": ["legajo", "nombre", "apellido", "rol"],
    "vehiculos": ["patente", "marca", "modelo", "anio"],
    "materiales": ["codigo", "nombre", "costo_unitario"],
    "stock": ["codigo_material", "fecha_ingreso", "cantidad"],
    "clientes": ["numero", "nombre"],
    "demanda": ["fecha", "codigo_tipo_tarea", "cantidad_clientes"],
    "ordenes": ["numero", "codigo_tipo_tarea", "fecha_programada"],
}


def leer_filas(archivo):
    nombre = archivo.name.lower()
    if nombre.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(archivo, read_only=True, data_only=True)
        filas = list(wb.active.iter_rows(values_only=True))
    else:
        contenido = archivo.read().decode("utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(contenido[:2000], delimiters=",;\t")
        filas = list(csv.reader(io.StringIO(contenido), dialecto))
    if not filas:
        return []
    cab = [texto(c, False).lower().replace(" ", "_") for c in filas[0]]
    return [dict(zip(cab, f)) for f in filas[1:] if any(c not in (None, "") for c in f)]


def importar(tipo, archivo):
    _, funcion, columnas, _, _ = IMPORTADORES[tipo]
    filas = leer_filas(archivo)
    errores, creados, actualizados = [], 0, 0
    if filas:
        faltan = [c for c in OBLIGATORIAS[tipo] if c not in filas[0]]
        if faltan:
            return {"errores": [(1, f"faltan columnas: {', '.join(faltan)}")], "creados": 0, "actualizados": 0}
    with transaction.atomic():
        for i, f in enumerate(filas, start=2):
            try:
                with transaction.atomic():
                    if funcion(f):
                        creados += 1
                    else:
                        actualizados += 1
            except ErrorFila as e:
                errores.append((i, str(e)))
            except Exception as e:  # noqa: BLE001 — se informa cualquier problema de la fila
                errores.append((i, f"error inesperado: {e}"))
        if errores:
            transaction.set_rollback(True)
    return {"errores": errores, "creados": creados, "actualizados": actualizados, "total": len(filas)}


@requiere_rol(GERENCIA)
def vista(request):
    resultado = None
    if request.method == "POST":
        tipo, archivo = request.POST.get("tipo"), request.FILES.get("archivo")
        if tipo in IMPORTADORES and archivo:
            resultado = importar(tipo, archivo)
            resultado["tipo"] = IMPORTADORES[tipo][0]
            if not resultado["errores"]:
                messages.success(request, f"Importación correcta: {resultado['creados']} nuevos, "
                                          f"{resultado['actualizados']} actualizados.")
                return redirect("tablero:importar")
        else:
            messages.error(request, "Elegí el tipo de datos y un archivo.")
    return render(request, "tablero/importar.html", {
        "importadores": [(k, v[0], v[2], v[4]) for k, v in IMPORTADORES.items()], "resultado": resultado})


@requiere_rol(GERENCIA)
def plantilla(request, tipo):
    _, _, columnas, ejemplo, _ = IMPORTADORES[tipo]
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="plantilla_{tipo}.csv"'
    resp.write("﻿")
    w = csv.writer(resp, delimiter=";")
    w.writerow(columnas)
    w.writerow(ejemplo)
    return resp
