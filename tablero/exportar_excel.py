"""Exportación del panel de productividad a Excel (.xlsx) con formato y gráficos nativos."""
from io import BytesIO
from statistics import mean

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.models import Indicador
from core.roles import GERENCIA, SUPERVISOR, persona_de, requiere_rol, rol_de

from .metricas import referencia_equipo, tableros_supervisores, tableros_tecnicos

AZUL = "2A78D6"
VERDE, AMARILLO, ROJO = "C8F2C8", "FFF1C2", "F8D0D0"
TITULO = Font(bold=True, size=14)
CABECERA = Font(bold=True, color="FFFFFF")
FONDO_CAB = PatternFill("solid", fgColor=AZUL)
BORDE = Border(bottom=Side(style="thin", color="E1E0D9"))


def _cabecera(ws, fila, columnas):
    for c, texto in enumerate(columnas, 1):
        celda = ws.cell(row=fila, column=c, value=texto)
        celda.font, celda.fill = CABECERA, FONDO_CAB
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[fila].height = 32


def _anchos(ws, anchos):
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a


def _semaforo_indice(ws, rango):
    """Colores del índice 0-100: verde ≥ 75, amarillo 55-74, rojo < 55."""
    ws.conditional_formatting.add(rango, CellIsRule(operator="greaterThanOrEqual", formula=["75"],
                                                    fill=PatternFill("solid", fgColor=VERDE)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="between", formula=["55", "74.999"],
                                                    fill=PatternFill("solid", fgColor=AMARILLO)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="lessThan", formula=["55"],
                                                    fill=PatternFill("solid", fgColor=ROJO)))


def _relleno_estado(estado):
    return PatternFill("solid", fgColor={"ok": VERDE, "aviso": AMARILLO, "critico": ROJO}.get(estado, "FFFFFF"))


def _r(v, n=1):
    return round(v, n) if v is not None else None


def _hoja_personas(wb, titulo, tableros, indicadores, sigla, extra_cols):
    ws = wb.create_sheet(titulo)
    cols = ["Legajo", "Apellido y nombre"] + [c for c, _ in extra_cols] + [sigla] + [
        f"{i.nombre} ({i.unidad})" for i in indicadores]
    ws.cell(row=1, column=1, value=f"{titulo} — valores y semáforo por indicador").font = TITULO
    meta_fila = ["", "Meta →"] + [""] * len(extra_cols) + [""] + [
        ("≤ " if not i.mayor_es_mejor else "") + f"{float(i.meta):g}" for i in indicadores]
    _cabecera(ws, 3, cols)
    for c, v in enumerate(meta_fila, 1):
        ws.cell(row=4, column=c, value=v).font = Font(italic=True, color="898781")
    for f, tb in enumerate(tableros, 5):
        p = tb.persona
        valores = [p.legajo, f"{p.apellido}, {p.nombre}"] + [fn(tb) for _, fn in extra_cols] + [_r(tb.indice)]
        for c, v in enumerate(valores, 1):
            ws.cell(row=f, column=c, value=v).border = BORDE
        for k, i in enumerate(indicadores):
            m = tb.get(i.codigo)
            celda = ws.cell(row=f, column=len(valores) + 1 + k, value=_r(m.valor) if m else None)
            celda.border = BORDE
            if m and m.valor is not None:
                celda.fill = _relleno_estado(m.estado)
    col_ind = 3 + len(extra_cols)
    letra = get_column_letter(col_ind)
    if tableros:
        _semaforo_indice(ws, f"{letra}5:{letra}{4 + len(tableros)}")
    _anchos(ws, [10, 28] + [16] * len(extra_cols) + [8] + [14] * len(indicadores))
    ws.freeze_panes = ws.cell(row=5, column=col_ind + 1)
    return ws


def construir_libro(dias=30, tecnicos_qs=None, solo_tecnicos=False):
    hoy = timezone.localdate()
    tt = tableros_tecnicos(dias=dias, tecnicos=tecnicos_qs)
    ts = [] if solo_tecnicos else tableros_supervisores(dias=dias)
    ind_t = list(Indicador.objects.filter(rol="tecnico", activo=True))
    ind_s = list(Indicador.objects.filter(rol="supervisor", activo=True))
    wb = Workbook()

    # ---- Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Virguel — Panel de productividad"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Período: últimos {dias} días hasta el {hoy:%d/%m/%Y} · generado {timezone.localtime():%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(color="52514E")
    con = [t for t in tt if t.indice is not None]
    con_igs = [t.indice for t in ts if t.indice is not None]
    filas = [
        ("IPT promedio (técnicos)", _r(mean(t.indice for t in con)) if con else None),
        ("IGS promedio (supervisores)", _r(mean(con_igs)) if con_igs else None),
        ("Técnicos medidos", len(con)),
        ("Técnicos con IPT ≥ 75 (bien)", sum(1 for t in con if t.estado == "ok")),
        ("Técnicos con IPT 55–74 (intermedio)", sum(1 for t in con if t.estado == "aviso")),
        ("Técnicos con IPT < 55 (bajo)", sum(1 for t in con if t.estado == "critico")),
    ]
    _cabecera(ws, 4, ["Indicador", "Valor"])
    for f, (k, v) in enumerate(filas, 5):
        ws.cell(row=f, column=1, value=k)
        ws.cell(row=f, column=2, value=v)
    ws["A12"] = "Cómo leerlo"
    ws["A12"].font = Font(bold=True)
    notas = ["IPT = Índice de Productividad del Técnico; IGS = Índice de Gestión del Supervisor (0–100).",
             "Cada indicador vale 100 puntos en la meta y 0 en el mínimo aceptable; el índice es el promedio ponderado.",
             "Verde ≥ 75 · amarillo 55–74 · rojo < 55. Definiciones completas en docs/METRICAS.md."]
    for k, n in enumerate(notas, 13):
        ws.cell(row=k, column=1, value=n)
    _anchos(ws, [44, 14])

    # ---- Indicadores de la empresa (mediana vs. meta)
    ws = wb.create_sheet("Indicadores empresa")
    ws["A1"] = "Valor típico del plantel (mediana) contra la meta"
    ws["A1"].font = TITULO
    _cabecera(ws, 3, ["Rol", "Indicador", "Unidad", "Valor empresa", "Meta", "Mínimo", "Puntos (0-100)", "Peso",
                      "Qué mide"])
    f = 4
    for rol, tableros, inds in (("Técnicos", tt, ind_t), ("Supervisores", ts, ind_s)):
        if not tableros:
            continue
        ref = referencia_equipo(tableros)
        for i in inds:
            v = ref.get(i.codigo)
            pts = i.puntos(v)
            fila = [rol, i.nombre, i.unidad, _r(v), float(i.meta), float(i.minimo), _r(pts, 0), i.peso, i.descripcion]
            for c, x in enumerate(fila, 1):
                ws.cell(row=f, column=c, value=x).border = BORDE
            if pts is not None:
                ws.cell(row=f, column=7).fill = _relleno_estado("ok" if pts >= 80 else "aviso" if pts >= 50 else "critico")
            f += 1
    _anchos(ws, [13, 36, 15, 14, 9, 9, 13, 7, 90])
    ws.freeze_panes = "C4"

    # ---- Técnicos y supervisores (detalle completo)
    _hoja_personas(wb, "Técnicos", tt, ind_t, "IPT", [
        ("Supervisor", lambda tb: tb.persona.supervisor.apellido if tb.persona.supervisor else ""),
        ("Zona", lambda tb: tb.persona.zona.nombre if tb.persona.zona else "")])
    if ts:
        _hoja_personas(wb, "Supervisores", ts, ind_s, "IGS", [])

    # ---- Por equipo y por zona
    from collections import defaultdict
    por_sup, por_zona = defaultdict(list), defaultdict(list)
    for t in con:
        por_sup[t.persona.supervisor.apellido if t.persona.supervisor else "Sin supervisor"].append(t)
        por_zona[t.persona.zona.nombre if t.persona.zona else "Sin zona"].append(t)
    igs = {t.persona.apellido: t.indice for t in ts}
    ws = wb.create_sheet("Por equipo")
    ws["A1"] = "Productividad por equipo"
    ws["A1"].font = TITULO
    _cabecera(ws, 3, ["Supervisor", "Técnicos", "IPT promedio del equipo", "IGS del supervisor",
                      "Eficiencia promedio (%)", "Primera visita (%)"])
    equipos = sorted(por_sup.items(), key=lambda kv: -mean(t.indice for t in kv[1]))
    for f, (s, lista) in enumerate(equipos, 4):
        def prom(cod):
            xs = [t.get(cod).valor for t in lista if t.get(cod) and t.get(cod).valor is not None]
            return _r(mean(xs)) if xs else None
        for c, x in enumerate([s, len(lista), _r(mean(t.indice for t in lista)), _r(igs.get(s)),
                               prom("eficiencia_jornada"), prom("primera_visita")], 1):
            ws.cell(row=f, column=c, value=x).border = BORDE
    n_eq = len(equipos)
    if n_eq:
        _semaforo_indice(ws, f"C4:D{3 + n_eq}")
    _anchos(ws, [22, 10, 16, 16, 16, 16])
    ws_z = wb.create_sheet("Por zona")
    ws_z["A1"] = "Productividad por zona"
    ws_z["A1"].font = TITULO
    _cabecera(ws_z, 3, ["Zona", "Técnicos", "IPT promedio"])
    zonas = sorted(por_zona.items(), key=lambda kv: -mean(t.indice for t in kv[1]))
    for f, (z, lista) in enumerate(zonas, 4):
        for c, x in enumerate([z, len(lista), _r(mean(t.indice for t in lista))], 1):
            ws_z.cell(row=f, column=c, value=x)
    if zonas:
        _semaforo_indice(ws_z, f"C4:C{3 + len(zonas)}")
    _anchos(ws_z, [20, 10, 14])

    # ---- Evolución
    ws_e = None
    if not solo_tecnicos:
        from .productividad import _evolucion
        ws_e = wb.create_sheet("Evolución")
        ws_e["A1"] = "Evolución semanal de los índices (promedio de la empresa)"
        ws_e["A1"].font = TITULO
        _cabecera(ws_e, 3, ["Semana", "IPT promedio", "IGS promedio"])
        evol = _evolucion(hoy, dias)
        for f, (fecha, a, b) in enumerate(evol, 4):
            ws_e.cell(row=f, column=1, value=fecha.strftime("%d/%m/%Y"))
            ws_e.cell(row=f, column=2, value=_r(a))
            ws_e.cell(row=f, column=3, value=_r(b))
        _anchos(ws_e, [14, 14, 14])

    # ---- Gráficos (nativos de Excel, editables)
    wg = wb.create_sheet("Gráficos", 1)
    wg["A1"] = "Gráficos del panel (se actualizan si se editan los datos de las otras hojas)"
    wg["A1"].font = TITULO
    hoja_t = wb["Técnicos"]
    n_t = len(tt)
    if n_t:
        ch = BarChart()
        ch.type, ch.style = "bar", 10
        ch.title, ch.y_axis.title = "IPT por técnico", "IPT (0-100)"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
        ch.add_data(Reference(hoja_t, min_col=5, min_row=5, max_row=4 + n_t), titles_from_data=False)
        ch.set_categories(Reference(hoja_t, min_col=2, min_row=5, max_row=4 + n_t))
        ch.legend = None
        ch.gapWidth = 40
        ch.x_axis.scaling.orientation = "maxMin"  # el mejor arriba, como en la planilla
        ch.height, ch.width = max(8, n_t * 0.45), 18
        wg.add_chart(ch, "A3")
    if n_eq:
        ws_eq = wb["Por equipo"]
        ch = BarChart()
        ch.title, ch.y_axis.title = "Por equipo: IPT del equipo vs. IGS del supervisor", "0-100"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
        ch.add_data(Reference(ws_eq, min_col=3, max_col=4, min_row=3, max_row=3 + n_eq), titles_from_data=True)
        ch.set_categories(Reference(ws_eq, min_col=1, min_row=4, max_row=3 + n_eq))
        ch.gapWidth = 60
        ch.height, ch.width = 8, 18
        wg.add_chart(ch, "L3")
    if ws_e is not None and ws_e.max_row > 3:
        ch = LineChart()
        ch.title, ch.y_axis.title = "Evolución de los índices", "0-100"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
        ch.add_data(Reference(ws_e, min_col=2, max_col=3, min_row=3, max_row=ws_e.max_row), titles_from_data=True)
        ch.set_categories(Reference(ws_e, min_col=1, min_row=4, max_row=ws_e.max_row))
        ch.height, ch.width = 8, 18
        wg.add_chart(ch, "L21")
    if len(zonas):
        ch = BarChart()
        ch.title = "IPT promedio por zona"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
        ch.add_data(Reference(ws_z, min_col=3, min_row=3, max_row=3 + len(zonas)), titles_from_data=True)
        ch.set_categories(Reference(ws_z, min_col=1, min_row=4, max_row=3 + len(zonas)))
        ch.gapWidth = 60
        ch.legend = None
        ch.height, ch.width = 8, 18
        wg.add_chart(ch, "L39")
    return wb


@requiere_rol(GERENCIA, SUPERVISOR)
def exportar(request):
    try:
        dias = max(7, min(180, int(request.GET.get("dias", 30))))
    except ValueError:
        dias = 30
    if rol_de(request.user) == SUPERVISOR:
        from core.models import Persona
        wb = construir_libro(dias, Persona.objects.filter(rol="tecnico", activo=True,
                                                          supervisor=persona_de(request.user)), solo_tecnicos=True)
    else:
        wb = construir_libro(dias)
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="productividad_{timezone.localdate():%Y%m%d}_{dias}d.xlsx"'
    return resp
